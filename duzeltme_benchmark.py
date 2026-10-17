"""OCR sonrası düzeltme katmanını sızıntısız ölçer.

Referans (Excel geçmişi + kullanıcı düzeltmeleri), değerlendirilen fişlerin
Excel satırları ve düzeltmeleri ÇIKARILARAK yeniden kurulur. Böylece düzeltme
katmanı cevabı ezberden bulamaz.

- Ayar: test seti dışındaki onaylı fişlerin kayıtlı ham OCR çıktıları, 2 katlı.
- Test: sabit 100 fişlik setin taze OCR çıktıları (ocr_benchmark --set raporu).

    python duzeltme_benchmark.py --excel "..\\NUSRET MAZOT 2026_YENİ 2.xlsx" \
        --rapor benchmark_set_qwen25vl7b_1600.json --output benchmark_duzeltme.json
"""

from __future__ import annotations

import argparse
import copy
import json
import itertools
import tempfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

import excel_ogrenme
import otomatik_duzeltme
from excel_ogrenme import ALAN_SUTUNLARI, _anlamli, arama_anahtari
from extract_json import dogrula, esle
from normalizasyon import firmayi_duzelt
from ocr_benchmark import FIELDS, ROOT, normalize, onayli_kayitlar, ozet, puanla

METIN_ALANLARI = tuple(ALAN_SUTUNLARI)


def excel_satirlari(excel: Path) -> list[dict]:
    wb = load_workbook(excel, read_only=True, data_only=True)
    satirlar = []
    try:
        for sayfa in wb.sheetnames:
            ws = wb[sayfa]
            try:
                excel_ogrenme._basliklari_dogrula(ws)
            except ValueError:
                continue
            for no, satir in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
                degerler = {
                    alan: (satir[sutun - 1] if len(satir) >= sutun else None)
                    for alan, sutun in ALAN_SUTUNLARI.items()
                }
                if any(_anlamli(v, a) for a, v in degerler.items()):
                    satirlar.append({"sayfa": sayfa, "satir": no, **degerler})
    finally:
        wb.close()
    return satirlar


def ocr_ciktisi(kayit: dict) -> dict:
    """Kayıtlı ham model çıktısını uygulamanın son işlemesinden geçirir."""
    return dogrula(esle(copy.deepcopy(kayit["_ham"])))


class Referans:
    """Belirli fişler hariç tutularak kurulmuş geçici referans dosyaları."""

    def __init__(self, excel: Path, satirlar: list[dict], kayitlar: list[dict],
                 haric: set[str]):
        self.dizin = Path(tempfile.mkdtemp(prefix="duzeltme_ref_"))
        haric_satir = {
            (k.get("_excel_sayfa"), k.get("_excel_satir"))
            for k in kayitlar
            if k.get("_dosya_hash") in haric
            and Path(str(k.get("_excel_dosya") or "")).name == excel.name
        }
        self.satirlar = [
            s for s in satirlar if (s["sayfa"], s["satir"]) not in haric_satir
        ]
        self.kayitlar = [
            k for k in kayitlar
            if k.get("_dosya_hash") not in haric and isinstance(k.get("_ham"), dict)
        ]
        self.haric_satir_sayisi = len(haric_satir)
        excel_ogrenme.REFERANS_FILE = self.dizin / "excel_referans.json"
        excel_ogrenme.KULLANICI_FILE = self.dizin / "kullanici_ogrenmeleri.json"
        self._excel_yaz(excel)
        excel_ogrenme.referans_yukle.cache_clear()
        excel_ogrenme.birlesik_referans_yukle.cache_clear()
        for kayit in self.kayitlar:
            onceki = ocr_ciktisi(kayit)
            degisen = {
                a for a in METIN_ALANLARI
                if arama_anahtari(onceki.get(a)) != arama_anahtari(kayit.get(a))
            }
            excel_ogrenme.duzeltmelerden_ogren(kayit, degisen, onceki)
        excel_ogrenme.birlesik_referans_yukle.cache_clear()
        ornekler = self.kombinasyonlar()
        excel_ogrenme.birliktelik_ornekleri = lambda: (ornekler, id(self))

    def _excel_yaz(self, excel: Path) -> None:
        kaynaklar = []
        for sayfa in sorted({s["sayfa"] for s in self.satirlar}):
            sayaclar = {a: Counter() for a in METIN_ALANLARI}
            for s in self.satirlar:
                if s["sayfa"] != sayfa:
                    continue
                for a in METIN_ALANLARI:
                    if _anlamli(s[a], a):
                        sayaclar[a][str(s[a]).strip()] += 1
            alanlar = {}
            for a, sayac in sayaclar.items():
                gruplar: dict[str, Counter] = {}
                for deger, adet in sayac.items():
                    gruplar.setdefault(arama_anahtari(deger), Counter())[deger] += adet
                alanlar[a] = [
                    {"deger": v.most_common(1)[0][0], "adet": sum(v.values())}
                    for v in gruplar.values()
                ]
            kaynaklar.append({"kaynak": str(excel), "kaynak_adi": excel.name,
                              "sayfa": sayfa, "alanlar": alanlar})
        excel_ogrenme.REFERANS_FILE.write_text(
            json.dumps({"surum": 2, "kaynaklar": kaynaklar}, ensure_ascii=False),
            encoding="utf-8",
        )

    def kombinasyonlar(self) -> list[dict]:
        """Geçmişte birlikte görülen (firma, plaka, ad_soyad) üçlüleri."""
        ornekler = [
            {a: s.get(a) for a in METIN_ALANLARI} for s in self.satirlar
        ] + [
            {a: k.get(a) for a in METIN_ALANLARI} for k in self.kayitlar
        ]
        return ornekler


YONTEMLER = {}


def yontem(ad):
    def kaydet(fonk):
        YONTEMLER[ad] = fonk
        return fonk
    return kaydet


@yontem("yok")
def _yok(kayit: dict, ref: Referans) -> dict:
    return kayit


@yontem("uygulama")
def _uygulama(kayit: dict, ref: Referans) -> dict:
    firma_temizle(kayit)
    excel_ogrenme.guvenli_otomatik_duzelt(kayit)
    return kayit


def _sozluk(ref: Referans):
    if getattr(ref, "_sozluk", None) is None:
        birlesik = excel_ogrenme.birlesik_referans_yukle() or {"alanlar": {}}
        ref._sozluk = otomatik_duzeltme.Sozluk(birlesik["alanlar"], ref.kombinasyonlar())
    return ref._sozluk


def firma_temizle(kayit: dict) -> None:
    """Tahminler eski normalizasyonla üretildiyse firmayı güncel kuralla yeniden düzelt."""
    kayit["firma"] = firmayi_duzelt(kayit.get("firma"))


def _kullanici_kurali(kayit: dict) -> set[str]:
    kullanici = excel_ogrenme._kullanici_yukle()
    uygulanan = set()
    for alan in METIN_ALANLARI:
        mevcut = str(kayit.get(alan) or "").strip()
        if not _anlamli(mevcut, alan):
            continue
        kesin = excel_ogrenme._kesin_kullanici_eslesmesi(kullanici, alan, mevcut)
        if kesin and str(kesin.get("dogru") or "").strip():
            kayit[alan] = str(kesin["dogru"]).strip()
            uygulanan.add(alan)
    return uygulanan


AKTIF_AYAR = otomatik_duzeltme.Ayar()


@yontem("yeni")
def _yeni(kayit: dict, ref: Referans, ayar=None) -> dict:
    """Deney sürümü (uygulama ile aynı mantık, eşik taraması için)."""
    firma_temizle(kayit)
    _kullanici_kurali(kayit)
    otomatik_duzeltme.duzelt(kayit, _sozluk(ref), ayar or AKTIF_AYAR, anlamli=_anlamli)
    return kayit


@yontem("yeni_birliktesiz")
def _yeni_birliktesiz(kayit: dict, ref: Referans) -> dict:
    ayar = copy.deepcopy(AKTIF_AYAR)
    ayar.ortak_bonus = 0
    return _yeni(kayit, ref, ayar)


def degerlendir(ornekler: list[tuple[dict, dict]], ref: Referans, yontemler) -> dict:
    """ornekler: (test-seti biçiminde fiş, OCR tahmini) çiftleri."""
    sonuc = {}
    for ad in yontemler:
        sonuclar, duzelen, bozulan = [], Counter(), Counter()
        for fis, tahmin in ornekler:
            once = copy.deepcopy(tahmin)
            sonra = YONTEMLER[ad](copy.deepcopy(tahmin), ref)
            for a in FIELDS:
                once_ok = normalize(a, once.get(a)) == normalize(a, fis["etiket"].get(a))
                sonra_ok = normalize(a, sonra.get(a)) == normalize(a, fis["etiket"].get(a))
                duzelen[a] += (not once_ok) and sonra_ok
                bozulan[a] += once_ok and not sonra_ok
            sonuclar.append(puanla(fis, sonra))
        o = ozet(sonuclar)
        sonuc[ad] = {
            "ornek": o["ornek"],
            "alan_dogrulugu": o["alan_dogrulugu"],
            "tam_dogru_fis": o["tam_dogru_fis"],
            **{a: o["alan_bazinda"][a] for a in METIN_ALANLARI},
            "duzelen": {a: duzelen[a] for a in METIN_ALANLARI if duzelen[a]},
            "bozulan": {a: bozulan[a] for a in METIN_ALANLARI if bozulan[a]},
            "_sonuclar": sonuclar,
        }
    return sonuc


def fis_bicimi(kayit: dict) -> dict:
    return {"id": (kayit.get("_dosya_hash") or "")[:16], "dosya": kayit.get("_dosya"),
            "ekipman": bool(kayit.get("_ekipman_durumu")),
            "etiket": {a: kayit.get(a) for a in FIELDS}}


def tablo(baslik: str, sonuc: dict) -> None:
    print(f"\n== {baslik}")
    print(f"{'yöntem':12}{'n':>5}{'alan%':>8}{'tam%':>7}{'firma':>7}{'plaka':>7}"
          f"{'ad':>7}  düzelen / bozulan")
    for ad, s in sonuc.items():
        print(f"{ad:12}{s['ornek']:>5}{s['alan_dogrulugu']:>8}{s['tam_dogru_fis']:>7}"
              f"{s['firma']:>7}{s['plaka']:>7}{s['ad_soyad']:>7}  {s['duzelen']} / {s['bozulan']}")


def izgara(excel, satirlar, kayitlar, ayar, test_hash) -> None:
    katlar = [ayar[0::2], ayar[1::2]]
    hazir = []
    for kat in katlar:
        ref = Referans(excel, satirlar, kayitlar, test_hash | {k["_dosya_hash"] for k in kat})
        _sozluk(ref)
        ornekler = [(fis_bicimi(k), ocr_ciktisi(k)) for k in kat]
        # Referans dosyaları her katta değiştiği için kullanıcı kuralını önceden uygula
        on = []
        for fis, t in ornekler:
            t = copy.deepcopy(t)
            firma_temizle(t)
            _kullanici_kurali(t)
            on.append((fis, t))
        hazir.append((ref, on))
    sonuclar = []
    for ef, ep, ea, fark, bonus, adet in itertools.product(
        (80, 85, 90), (75, 80, 85, 90), (75, 80, 85, 90), (3, 5, 8), (0, 5, 10), (1, 2)
    ):
        ay = otomatik_duzeltme.Ayar({"firma": ef, "plaka": ep, "ad_soyad": ea}, fark, bonus, adet)
        puanlar = []
        for ref, on in hazir:
            for fis, t in on:
                t = copy.deepcopy(t)
                otomatik_duzeltme.duzelt(t, ref._sozluk, ay, anlamli=_anlamli)
                puanlar.append(puanla(fis, t))
        o = ozet(puanlar)
        sonuclar.append((o["tam_dogru_fis"], o["alan_dogrulugu"], ef, ep, ea, fark, bonus, adet,
                         o["alan_bazinda"]["firma"], o["alan_bazinda"]["plaka"],
                         o["alan_bazinda"]["ad_soyad"]))
    sonuclar.sort(reverse=True)
    print("tam% alan% | esik firma/plaka/ad fark bonus adet | firma plaka ad")
    for r in sonuclar[:15]:
        print(r)


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--excel", type=Path, required=True)
    p.add_argument("--set", type=Path, default=ROOT / "benchmark_testseti.json")
    p.add_argument("--rapor", type=Path, help="ocr_benchmark --set çıktısı (test)")
    p.add_argument("--yontem", nargs="*", default=None)
    p.add_argument("--sadece", choices=["ayar", "test"], default=None)
    p.add_argument("--output", type=Path)
    p.add_argument("--izgara", action="store_true", help="ayar setinde eşik taraması")
    a = p.parse_args()
    yontemler = a.yontem or list(YONTEMLER)

    satirlar = excel_satirlari(a.excel)
    kayitlar = [k for k in onayli_kayitlar() if isinstance(k.get("_ham"), dict)]
    test_set = json.loads(a.set.read_text(encoding="utf-8"))
    test_hash = {f["sha256"] for f in test_set["fisler"]}
    ayar = sorted((k for k in kayitlar if k.get("_dosya_hash") not in test_hash),
                  key=lambda k: k["_dosya_hash"])
    rapor = {}

    if a.izgara:
        izgara(a.excel, satirlar, kayitlar, ayar, test_hash)
        return

    if a.sadece in (None, "ayar"):
        katlar = [ayar[0::2], ayar[1::2]]
        birlesik = {}
        for kat in katlar:
            ref = Referans(a.excel, satirlar, kayitlar,
                           test_hash | {k["_dosya_hash"] for k in kat})
            s = degerlendir([(fis_bicimi(k), ocr_ciktisi(k)) for k in kat], ref, yontemler)
            for ad, d in s.items():
                birlesik.setdefault(ad, []).extend(d["_sonuclar"])
        ayar_sonuc = {}
        for ad, sonuclar in birlesik.items():
            o = ozet(sonuclar)
            ayar_sonuc[ad] = {"ornek": o["ornek"], "alan_dogrulugu": o["alan_dogrulugu"],
                              "tam_dogru_fis": o["tam_dogru_fis"],
                              **{x: o["alan_bazinda"][x] for x in METIN_ALANLARI},
                              "duzelen": {}, "bozulan": {}}
        tablo(f"AYAR (test dışı {len(ayar)} fiş, kayıtlı eski OCR, 2 kat)", ayar_sonuc)
        rapor["ayar"] = ayar_sonuc

    if a.sadece in (None, "test") and a.rapor:
        ocr = {r["id"]: r for r in json.loads(a.rapor.read_text(encoding="utf-8"))["sonuclar"]}
        ornekler = []
        for fis in test_set["fisler"]:
            r = ocr[fis["id"]]
            tahmin = {x: d["okunan"] for x, d in r["alanlar"].items()}
            tahmin["_ekipman_durumu"] = fis["ekipman"]
            ornekler.append((fis, tahmin))
        ref = Referans(a.excel, satirlar, kayitlar, test_hash)
        print(f"\nTest referansı: {len(ref.satirlar)} Excel satırı "
              f"({ref.haric_satir_sayisi} test satırı çıkarıldı), {len(ref.kayitlar)} onaylı fiş")
        test_sonuc = degerlendir(ornekler, ref, yontemler)
        tablo(f"TEST (sabit set, {a.rapor.name})", test_sonuc)
        rapor["test"] = test_sonuc

    if a.output:
        a.output.write_text(json.dumps(rapor, ensure_ascii=False, indent=1), encoding="utf-8")


if __name__ == "__main__":
    main()
