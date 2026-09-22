"""OCR normalizasyonu, öğrenme ve Excel aktarımı için güvenli entegrasyon testleri."""

from __future__ import annotations

import hashlib
import json
import shutil
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import excel_aktar
import excel_ogrenme
import extract_json
import uygulama
import parti_yonetimi
from dogrulama import kaydi_dogrula
from normalizasyon import fis_normalize_et, kilometreyi_duzelt, plakayi_duzelt


ROOT = Path(__file__).resolve().parent
KAYNAK_EXCEL = Path.home() / "Documents" / "NUSRET MAZOT 2026_YENİ 2.xlsx"


def sha256(yol: Path) -> str:
    ozet = hashlib.sha256()
    with yol.open("rb") as dosya:
        for parca in iter(lambda: dosya.read(1024 * 1024), b""):
            ozet.update(parca)
    return ozet.hexdigest()


class NormalizasyonTesti(unittest.TestCase):
    def test_tek_haneli_kilometre_cizgi_olur(self):
        for deger in (0, 1, 7, 9, " 7 "):
            self.assertEqual(kilometreyi_duzelt(deger), "-")
        for deger in ("10", "123456", "-"):
            self.assertEqual(kilometreyi_duzelt(deger), deger)
        self.assertEqual(fis_normalize_et({"km": "1"})["km"], "-")

    def test_bosluk_firma_ve_plaka_kurallari(self):
        self.assertEqual(fis_normalize_et({"firma": "7"})["firma"], "Selim İnşaat")
        self.assertEqual(plakayi_duzelt("H2 1663"), "34 H2 1663")
        self.assertEqual(plakayi_duzelt("KAZIK MAK", ekipman_durumu=True), "KAZIK MAK")


class DogrulamaTesti(unittest.TestCase):
    def test_gecerli_arac_kaydi(self):
        kayit = kaydi_dogrula(fis_normalize_et({
            "fis_no": "001234", "tarih": "10/09/2026", "firma": "Selim İnşaat",
            "plaka": "34 ABC 123", "ad_soyad": "Ali Veli", "km": "123456",
            "litre": "250,5",
        }))
        self.assertFalse(kayit["_kontrol_gerekli"])

    def test_tum_kritik_hatalar_aciklanir(self):
        kayit = kaydi_dogrula(fis_normalize_et({
            "fis_no": "-12", "tarih": "01/01/2025", "firma": "1",
            "plaka": "34 A2 1", "ad_soyad": "-", "km": "7", "litre": "-",
        }))
        neden = kayit["_kontrol_nedeni"]
        for parca in ("fiş no", "tarih", "plaka", "ad soyad", "kilometre", "litre"):
            self.assertIn(parca, neden)


class OtomatikYedekTesti(unittest.TestCase):
    def test_gunluk_json_yedegi_bir_kez_olusturulur(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as gecici:
            root = Path(gecici)
            (root / "sonuclar.json").write_text("[]", encoding="utf-8")
            ilk = uygulama.gunluk_veri_yedegi(root)
            ikinci = uygulama.gunluk_veri_yedegi(root)
            self.assertIsNotNone(ilk)
            self.assertEqual(ilk, ikinci)
            self.assertTrue(ilk.is_file())


class PartiYonetimiTesti(unittest.TestCase):
    def test_yeni_parti_eski_veriyi_arsivler_ve_aktif_listeyi_temizler(self):
        eski = (
            parti_yonetimi.ROOT, parti_yonetimi.DATA_FILE,
            parti_yonetimi.FISLER_DIR,
        )
        with tempfile.TemporaryDirectory(dir=ROOT) as gecici:
            root = Path(gecici)
            data = root / "sonuclar.json"
            fisler = root / "fisler"
            fisler.mkdir()
            data.write_text(json.dumps([
                {"_dosya": "001.jpg", "_onaylandi": True, "_excel_aktarildi": True}
            ]), encoding="utf-8")
            (fisler / "001.jpg").write_bytes(b"ornek")
            try:
                parti_yonetimi.yollari_ayarla(root, data, fisler)
                sonuc = parti_yonetimi.yeni_parti(
                    "İkinci Parti", {"aktarilan": 1, "sayfa": "2026"}
                )
                self.assertEqual(sonuc["aktif"]["ad"], "İkinci Parti")
                self.assertEqual(json.loads(data.read_text(encoding="utf-8")), [])
                self.assertTrue(fisler.is_dir())
                arsiv = sonuc["arsivlenen"]
                arsiv_dir = parti_yonetimi.ARSIV_DIR / arsiv["id"]
                self.assertTrue((arsiv_dir / "fisler" / "001.jpg").is_file())
                manifest = json.loads((arsiv_dir / "manifest.json").read_text(encoding="utf-8"))
                self.assertEqual(manifest["kayit_sayisi"], 1)
                self.assertEqual(manifest["gorsel_sayisi"], 1)
                self.assertEqual(manifest["eksik_gorseller"], [])
                self.assertEqual(len(parti_yonetimi.partileri_listele()), 2)
            finally:
                parti_yonetimi.yollari_ayarla(*eski)


class EkipmanFisiTesti(unittest.TestCase):
    """Ekipman fişlerinde ekipman adı plakaya, kişi adı ad_soyada gider."""

    def _esle(self, arac, kisi, km="-"):
        return extract_json.esle({
            "arac_plakasi_alani": arac,
            "adi_soyadi_alani": kisi,
            "km_alani": km,
        })

    def test_ekipman_plaka_hucresindeyse_takas_edilmez(self):
        sonuc = self._esle("SİLİNDİR", "ATABEK")
        self.assertTrue(sonuc["_ekipman_durumu"])
        self.assertEqual(sonuc["plaka"], "SİLİNDİR")
        self.assertEqual(sonuc["ad_soyad"], "ATABEK")

    def test_ekipman_isim_hucresindeyse_takas_edilir(self):
        sonuc = self._esle("HASAN ZEMIN", "KAZIK MAK")
        self.assertEqual(sonuc["plaka"], "KAZIK MAK")
        self.assertEqual(sonuc["ad_soyad"], "HASAN ZEMIN")

    def test_ocr_rakama_cevirse_bile_ekipman_taninir(self):
        sonuc = self._esle("KA21K MAK", "HİSAN ZEMIN")
        self.assertTrue(sonuc["_ekipman_durumu"])
        self.assertEqual(sonuc["plaka"], "KA21K MAK")
        self.assertEqual(sonuc["ad_soyad"], "HİSAN ZEMIN")

    def test_gercek_plaka_ekipman_sayilmaz(self):
        sonuc = self._esle("34 DSN 621", "MURAT SARI", "120500")
        self.assertFalse(sonuc["_ekipman_durumu"])
        self.assertEqual(sonuc["plaka"], "34 DSN 621")
        self.assertEqual(sonuc["ad_soyad"], "MURAT SARI")

    def test_km_hucresindeki_isim_eklenir_doldurma_degeri_eklenmez(self):
        self.assertEqual(
            self._esle("HASAN ZEMIN", "KAZIK MAK", "Veli Telep")["ad_soyad"],
            "HASAN ZEMIN - Veli Telep",
        )
        self.assertEqual(self._esle("SİLİNDİR", "ATABEK", "K.M.")["ad_soyad"], "ATABEK")


class KullaniciOgrenmesiTesti(unittest.TestCase):
    def test_yalniz_degistirilen_metin_alanlari_ogrenilir(self):
        eski_dosya = excel_ogrenme.KULLANICI_FILE
        with tempfile.TemporaryDirectory(dir=ROOT) as gecici:
            try:
                excel_ogrenme.KULLANICI_FILE = Path(gecici) / "ogrenilen.json"
                excel_ogrenme.birlesik_referans_yukle.cache_clear()
                adet = excel_ogrenme.duzeltmelerden_ogren(
                    {"firma": "DOĞRU YENİ FİRMA", "km": "123456"},
                    {"firma", "km"},
                )
                self.assertEqual(adet, 1)
                oneri = excel_ogrenme.kayit_onerileri({
                    "firma": "DOGRU YENI FIRMAA", "plaka": "-", "ad_soyad": "-",
                })
                self.assertEqual(oneri["oneriler"][0]["onerilen"], "DOĞRU YENİ FİRMA")
                kayit = json.loads(excel_ogrenme.KULLANICI_FILE.read_text(encoding="utf-8"))
                self.assertNotIn("km", {
                    alan for alan, liste in kayit["alanlar"].items() if liste
                })
            finally:
                excel_ogrenme.KULLANICI_FILE = eski_dosya
                excel_ogrenme.birlesik_referans_yukle.cache_clear()

    def test_yuksek_guvenli_gecmis_degerleri_otomatik_duzeltir(self):
        eski_referans = excel_ogrenme.REFERANS_FILE
        eski_kullanici = excel_ogrenme.KULLANICI_FILE
        with tempfile.TemporaryDirectory(dir=ROOT) as gecici:
            try:
                gecici = Path(gecici)
                excel_ogrenme.REFERANS_FILE = gecici / "referans.json"
                excel_ogrenme.KULLANICI_FILE = gecici / "kullanici.json"
                excel_ogrenme.KULLANICI_FILE.write_text(json.dumps({
                    "surum": 1,
                    "alanlar": {
                        "firma": [],
                        "plaka": [{"deger": "34 DTV 263", "adet": 4}],
                        "ad_soyad": [{"deger": "SEZGİN KOCAGİL", "adet": 3}],
                    },
                }, ensure_ascii=False), encoding="utf-8")
                excel_ogrenme.referans_yukle.cache_clear()
                excel_ogrenme.birlesik_referans_yukle.cache_clear()

                kayit = {
                    "firma": "-", "plaka": "34 ATV 263",
                    "ad_soyad": "SEZGIN KOCAGIL",
                }
                adet = excel_ogrenme.guvenli_otomatik_duzelt(kayit)
                self.assertEqual(adet, 2)
                self.assertEqual(kayit["plaka"], "34 DTV 263")
                self.assertEqual(kayit["ad_soyad"], "SEZGİN KOCAGİL")
                self.assertTrue(kayit["_otomatik_duzeltildi"])
            finally:
                excel_ogrenme.REFERANS_FILE = eski_referans
                excel_ogrenme.KULLANICI_FILE = eski_kullanici
                excel_ogrenme.referans_yukle.cache_clear()
                excel_ogrenme.birlesik_referans_yukle.cache_clear()

    def test_fazladan_kelimeli_aday_otomatik_uygulanmaz(self):
        """'RT 650' plakası 'RT 650 SARI VİNÇ'e çevrilmemeli."""
        eski_referans = excel_ogrenme.REFERANS_FILE
        eski_kullanici = excel_ogrenme.KULLANICI_FILE
        with tempfile.TemporaryDirectory(dir=ROOT) as gecici:
            try:
                gecici = Path(gecici)
                excel_ogrenme.REFERANS_FILE = gecici / "referans.json"
                excel_ogrenme.KULLANICI_FILE = gecici / "kullanici.json"
                excel_ogrenme.KULLANICI_FILE.write_text(json.dumps({
                    "surum": 2,
                    "alanlar": {
                        "firma": [],
                        "plaka": [
                            {"deger": "RT 650 SARI VİNÇ", "adet": 22},
                            {"deger": "34 DTV 263", "adet": 4},
                        ],
                        "ad_soyad": [],
                    },
                    "eslesmeler": {"firma": [], "plaka": [], "ad_soyad": []},
                }, ensure_ascii=False), encoding="utf-8")
                excel_ogrenme.referans_yukle.cache_clear()
                excel_ogrenme.birlesik_referans_yukle.cache_clear()

                kayit = {"firma": "-", "plaka": "RT 650", "ad_soyad": "-"}
                self.assertEqual(excel_ogrenme.guvenli_otomatik_duzelt(kayit), 0)
                self.assertEqual(kayit["plaka"], "RT 650")

                # Harf düzeyindeki gerçek OCR hatası düzeltilmeye devam etmeli.
                kayit = {"firma": "-", "plaka": "34 ATV 263", "ad_soyad": "-"}
                self.assertEqual(excel_ogrenme.guvenli_otomatik_duzelt(kayit), 1)
                self.assertEqual(kayit["plaka"], "34 DTV 263")
            finally:
                excel_ogrenme.REFERANS_FILE = eski_referans
                excel_ogrenme.KULLANICI_FILE = eski_kullanici
                excel_ogrenme.referans_yukle.cache_clear()
                excel_ogrenme.birlesik_referans_yukle.cache_clear()

    def test_ocr_yanlisi_ilk_kullanici_duzeltmesinden_sonra_hatirlanir(self):
        eski_referans = excel_ogrenme.REFERANS_FILE
        eski_kullanici = excel_ogrenme.KULLANICI_FILE
        with tempfile.TemporaryDirectory(dir=ROOT) as gecici:
            try:
                gecici = Path(gecici)
                excel_ogrenme.REFERANS_FILE = gecici / "referans.json"
                excel_ogrenme.KULLANICI_FILE = gecici / "kullanici.json"
                excel_ogrenme.referans_yukle.cache_clear()
                excel_ogrenme.birlesik_referans_yukle.cache_clear()

                excel_ogrenme.duzeltmelerden_ogren(
                    {"ad_soyad": "GÖRKEM KARAYEL"},
                    {"ad_soyad"},
                    {"ad_soyad": "GÖRKEM KARAHİL"},
                )
                yeni = {
                    "firma": "-", "plaka": "-",
                    "ad_soyad": "GÖRKEM KARAHİL",
                }
                self.assertEqual(excel_ogrenme.guvenli_otomatik_duzelt(yeni), 1)
                self.assertEqual(yeni["ad_soyad"], "GÖRKEM KARAYEL")
                self.assertIn(
                    "kullanıcı düzeltmesi",
                    yeni["_otomatik_duzeltmeler"][0]["neden"],
                )
            finally:
                excel_ogrenme.REFERANS_FILE = eski_referans
                excel_ogrenme.KULLANICI_FILE = eski_kullanici
                excel_ogrenme.referans_yukle.cache_clear()
                excel_ogrenme.birlesik_referans_yukle.cache_clear()

    def test_gercek_firma_baska_firmaya_cevrilmez(self):
        """Kullanıcının çok kez onayladığı gerçek bir değer OCR hatası sayılmaz."""
        eski_referans = excel_ogrenme.REFERANS_FILE
        eski_kullanici = excel_ogrenme.KULLANICI_FILE
        with tempfile.TemporaryDirectory(dir=ROOT) as gecici:
            try:
                gecici = Path(gecici)
                excel_ogrenme.REFERANS_FILE = gecici / "referans.json"
                excel_ogrenme.KULLANICI_FILE = gecici / "kullanici.json"
                excel_ogrenme.referans_yukle.cache_clear()
                excel_ogrenme.birlesik_referans_yukle.cache_clear()

                # Selim İnşaat defalarca onaylanmış gerçek bir firma.
                for _ in range(3):
                    excel_ogrenme.duzeltmelerden_ogren(
                        {"firma": "Selim İnşaat"}, {"firma"}, None
                    )
                # Bir fişte kullanıcı Selim İnşaat yazan alanı AZGÜN YAPI yapsa bile
                # bu bir OCR kuralına dönüşmemeli.
                excel_ogrenme.duzeltmelerden_ogren(
                    {"firma": "AZGÜN YAPI"}, {"firma"}, {"firma": "Selim İnşaat"}
                )
                veri = json.loads(excel_ogrenme.KULLANICI_FILE.read_text(encoding="utf-8"))
                self.assertEqual(veri["eslesmeler"].get("firma", []), [])

                yeni_kayit = {"firma": "Selim İnşaat", "plaka": "-", "ad_soyad": "-"}
                self.assertEqual(excel_ogrenme.guvenli_otomatik_duzelt(yeni_kayit), 0)
                self.assertEqual(yeni_kayit["firma"], "Selim İnşaat")
            finally:
                excel_ogrenme.REFERANS_FILE = eski_referans
                excel_ogrenme.KULLANICI_FILE = eski_kullanici
                excel_ogrenme.referans_yukle.cache_clear()
                excel_ogrenme.birlesik_referans_yukle.cache_clear()

    def test_celiskili_kural_otomatik_uygulanmaz(self):
        """Aynı OCR metni için iki farklı doğru varsa tahmin yürütülmez."""
        eski_referans = excel_ogrenme.REFERANS_FILE
        eski_kullanici = excel_ogrenme.KULLANICI_FILE
        with tempfile.TemporaryDirectory(dir=ROOT) as gecici:
            try:
                gecici = Path(gecici)
                excel_ogrenme.REFERANS_FILE = gecici / "referans.json"
                excel_ogrenme.KULLANICI_FILE = gecici / "kullanici.json"
                excel_ogrenme.referans_yukle.cache_clear()
                excel_ogrenme.birlesik_referans_yukle.cache_clear()
                excel_ogrenme.KULLANICI_FILE.write_text(json.dumps({
                    "surum": 2,
                    "alanlar": {"firma": [], "plaka": [], "ad_soyad": []},
                    "eslesmeler": {"firma": [
                        {"yanlis": "A2GÜN YAPI", "dogru": "AZGÜN YAPI", "adet": 3},
                        {"yanlis": "A2GÜN YAPI", "dogru": "AYGÜN YAPI", "adet": 2},
                    ], "plaka": [], "ad_soyad": []},
                }, ensure_ascii=False), encoding="utf-8")
                excel_ogrenme.birlesik_referans_yukle.cache_clear()

                kayit = {"firma": "A2GÜN YAPI", "plaka": "-", "ad_soyad": "-"}
                self.assertEqual(excel_ogrenme.guvenli_otomatik_duzelt(kayit), 0)
                self.assertEqual(kayit["firma"], "A2GÜN YAPI")
            finally:
                excel_ogrenme.REFERANS_FILE = eski_referans
                excel_ogrenme.KULLANICI_FILE = eski_kullanici
                excel_ogrenme.referans_yukle.cache_clear()
                excel_ogrenme.birlesik_referans_yukle.cache_clear()

    def test_onayli_gecmis_duzeltme_bir_kez_ogrenilir(self):
        eski_referans = excel_ogrenme.REFERANS_FILE
        eski_kullanici = excel_ogrenme.KULLANICI_FILE
        with tempfile.TemporaryDirectory(dir=ROOT) as gecici:
            try:
                gecici = Path(gecici)
                excel_ogrenme.REFERANS_FILE = gecici / "referans.json"
                excel_ogrenme.KULLANICI_FILE = gecici / "kullanici.json"
                excel_ogrenme.referans_yukle.cache_clear()
                excel_ogrenme.birlesik_referans_yukle.cache_clear()
                kayitlar = [{
                    "firma": "Selim İnşaat", "plaka": "34 DSN 536",
                    "ad_soyad": "GÖRKEM KARAYEL", "_onaylandi": True,
                    "_duzeltildi": True, "_dosya_hash": "tekil-hash",
                    "_ham": {
                        "sayin_alani": None, "arac_plakasi_alani": "DSN 536",
                        "adi_soyadi_alani": "GÖRKEM KARAHİL",
                    },
                }]
                self.assertEqual(
                    excel_ogrenme.gecmis_duzeltmelerden_eslesme_ogren(kayitlar), 1
                )
                self.assertEqual(
                    excel_ogrenme.gecmis_duzeltmelerden_eslesme_ogren(kayitlar), 0
                )
                veri = json.loads(excel_ogrenme.KULLANICI_FILE.read_text(encoding="utf-8"))
                self.assertEqual(veri["eslesmeler"]["ad_soyad"][0]["adet"], 1)
            finally:
                excel_ogrenme.REFERANS_FILE = eski_referans
                excel_ogrenme.KULLANICI_FILE = eski_kullanici
                excel_ogrenme.referans_yukle.cache_clear()
                excel_ogrenme.birlesik_referans_yukle.cache_clear()


@unittest.skipUnless(KAYNAK_EXCEL.exists(), "Gerçek Excel örneği bu bilgisayarda yok")
class BirikimliExcelOgrenmesiTesti(unittest.TestCase):
    def test_farkli_sayfalar_birikir_ayni_sayfa_guncellenir(self):
        eski_referans = excel_ogrenme.REFERANS_FILE
        eski_kullanici = excel_ogrenme.KULLANICI_FILE
        with tempfile.TemporaryDirectory(dir=ROOT) as gecici:
            try:
                excel_ogrenme.REFERANS_FILE = Path(gecici) / "referans.json"
                excel_ogrenme.KULLANICI_FILE = Path(gecici) / "kullanici.json"
                excel_ogrenme.referans_yukle.cache_clear()
                excel_ogrenme.birlesik_referans_yukle.cache_clear()
                ilk = excel_ogrenme.referans_olustur(KAYNAK_EXCEL, "2026 MAZOT")
                ikinci = excel_ogrenme.referans_olustur(KAYNAK_EXCEL, "2025 MAZOT ")
                durum = excel_ogrenme.durum()
                self.assertEqual((ilk["toplam_kaynak"], ikinci["toplam_kaynak"]), (1, 2))
                self.assertEqual(durum["kaynak_sayisi"], 2)
                self.assertEqual(durum["satir_sayisi"], ilk["satir_sayisi"] + ikinci["satir_sayisi"])

                tekrar = excel_ogrenme.referans_olustur(KAYNAK_EXCEL, "2026 MAZOT")
                self.assertEqual(tekrar["toplam_kaynak"], 2)
                self.assertEqual(excel_ogrenme.durum()["kaynak_sayisi"], 2)
            finally:
                excel_ogrenme.REFERANS_FILE = eski_referans
                excel_ogrenme.KULLANICI_FILE = eski_kullanici
                excel_ogrenme.referans_yukle.cache_clear()
                excel_ogrenme.birlesik_referans_yukle.cache_clear()


@unittest.skipUnless(KAYNAK_EXCEL.exists(), "Gerçek Excel örneği bu bilgisayarda yok")
class ExcelAktarimTesti(unittest.TestCase):
    def test_yanlis_dosyaya_aktarilanlar_yeniden_aktarilabilir(self):
        """Aktarılmış fişler 'tumu' seçeneğiyle başka bir deftere yeniden gider."""
        eski_root = excel_aktar.ROOT
        eski_data = excel_aktar.DATA_FILE
        eski_ayar = excel_aktar.AYAR_FILE
        with tempfile.TemporaryDirectory(dir=ROOT) as gecici:
            gecici = Path(gecici)
            yanlis = gecici / "yanlis_defter.xlsx"
            dogru = gecici / "dogru_defter.xlsx"
            shutil.copy2(KAYNAK_EXCEL, yanlis)
            shutil.copy2(KAYNAK_EXCEL, dogru)
            data = gecici / "sonuclar.json"
            data.write_text(json.dumps([{
                "_dosya": "test-1.jpg", "_onaylandi": True,
                "tarih": "10/09/2026", "fis_no": "009876",
                "firma": "TEST FİRMA", "plaka": "34 ABC 123",
                "ad_soyad": "TEST KİŞİ", "km": "123456", "litre": "42,5",
            }], ensure_ascii=False), encoding="utf-8")
            try:
                excel_aktar.DATA_FILE = data
                excel_aktar.AYAR_FILE = gecici / "ayar.json"
                excel_aktar.ROOT = gecici

                self.assertEqual(excel_aktar.aktar(yanlis, "2026 MAZOT")["aktarilan"], 1)
                # İkinci kez aynı fiş normalde aktarılmaz.
                self.assertEqual(excel_aktar.aktar(dogru, "2026 MAZOT")["aktarilan"], 0)
                # 'tumu' ile doğru deftere yeniden aktarılır.
                sonuc = excel_aktar.aktar(dogru, "2026 MAZOT", tumu=True)
                self.assertEqual(sonuc["aktarilan"], 1)

                wb = load_workbook(dogru)
                ws = wb["2026 MAZOT"]
                self.assertEqual(ws.cell(sonuc["ilk_satir"], 4).value, "34 ABC 123")
                wb.close()

                kaydedilen = json.loads(data.read_text(encoding="utf-8"))
                self.assertEqual(kaydedilen[0]["_excel_dosya"], str(dogru))
            finally:
                excel_aktar.ROOT = eski_root
                excel_aktar.DATA_FILE = eski_data
                excel_aktar.AYAR_FILE = eski_ayar

    def test_dogru_satir_ve_sutunlara_yazar(self):
        kaynak_hash = sha256(KAYNAK_EXCEL)
        eski_root = excel_aktar.ROOT
        eski_data = excel_aktar.DATA_FILE
        eski_ayar = excel_aktar.AYAR_FILE
        with tempfile.TemporaryDirectory(dir=ROOT) as gecici:
            gecici = Path(gecici)
            kopya = gecici / "aktarim_testi.xlsx"
            data = gecici / "sonuclar.json"
            ayar = gecici / "ayar.json"
            shutil.copy2(KAYNAK_EXCEL, kopya)
            kayitlar = [
                {
                    "_dosya": "test-1.jpg", "_onaylandi": True,
                    "tarih": "10/09/2026", "fis_no": "009876",
                    "firma": "TEST FİRMA", "plaka": "34 ABC 123",
                    "ad_soyad": "TEST KİŞİ", "km": "123456", "litre": "42,5",
                },
                {
                    "_dosya": "test-2.jpg", "_onaylandi": True,
                    "tarih": "10/09/2026", "fis_no": "009877",
                    "firma": "TEST FİRMA", "plaka": "34 ABC 124",
                    "ad_soyad": "TEST KİŞİ 2", "km": "-", "litre": "50",
                },
            ]
            data.write_text(json.dumps(kayitlar, ensure_ascii=False), encoding="utf-8")
            try:
                excel_aktar.DATA_FILE = data
                excel_aktar.AYAR_FILE = ayar
                excel_aktar.ROOT = gecici
                sonuc = excel_aktar.aktar(kopya, "2026 MAZOT")
                self.assertNotIn("hata", sonuc)
                self.assertEqual((sonuc["aktarilan"], sonuc["ilk_satir"], sonuc["son_satir"]), (2, 1308, 1309))

                wb = load_workbook(kopya, data_only=False)
                ws = wb["2026 MAZOT"]
                self.assertTrue(all(ws.cell(1307, c).value is None for c in (1, 2, 3, 4, 5, 6, 8)))
                self.assertEqual(ws.cell(1308, 1).value, datetime(2026, 9, 10))
                self.assertEqual(ws.cell(1308, 2).value, 9876)
                self.assertEqual(ws.cell(1308, 3).value, "TEST FİRMA")
                self.assertEqual(ws.cell(1308, 4).value, "34 ABC 123")
                self.assertEqual(ws.cell(1308, 5).value, "TEST KİŞİ")
                self.assertEqual(ws.cell(1308, 6).value, 123456)
                self.assertEqual(ws.cell(1308, 8).value, 42.5)
                self.assertIsNone(ws.cell(1309, 6).value)
                self.assertTrue(all(ws.cell(1308, c).value is None for c in (7, 9, 10, 11, 12, 13)))
                wb.close()

                kaydedilen = json.loads(data.read_text(encoding="utf-8"))
                self.assertTrue(all(kayit.get("_excel_aktarildi") for kayit in kaydedilen))
                self.assertEqual([k["_excel_satir"] for k in kaydedilen], [1308, 1309])
                self.assertTrue(Path(sonuc["yedek"]).is_file())
            finally:
                excel_aktar.ROOT = eski_root
                excel_aktar.DATA_FILE = eski_data
                excel_aktar.AYAR_FILE = eski_ayar

        self.assertEqual(sha256(KAYNAK_EXCEL), kaynak_hash)


if __name__ == "__main__":
    unittest.main(verbosity=2)
