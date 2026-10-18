"""
sonuclar.json (extract_json.py çıktısı) dosyasını fisler.xlsx'e çevirir.
'_kontrol_gerekli' işaretli satırlar sarı renkle vurgulanır, böylece
300 satırın tamamı yerine sadece şüpheli olanları elle kontrol edersin.

Kullanım:
    python -m tools.export_excel [sonuclar.json] [fisler.xlsx]
"""

import json
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from mazot.dogrulama import kaydi_dogrula
from mazot.normalizasyon import fis_normalize_et

SUTUNLAR = [
    ("_dosya", "Dosya"),
    ("fis_no", "Fiş No (MRK)"),
    ("tarih", "Tarih"),
    ("firma", "Firma"),
    ("plaka", "Plaka / Ekipman"),
    ("ad_soyad", "Ad Soyad"),
    ("km", "K.M."),
    ("litre", "Litre"),
    ("_kontrol_gerekli", "Kontrol Gerekli"),
    ("_kontrol_nedeni", "Kontrol Nedeni"),
]

SARI = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")


def excel_degeri(deger):
    """Excel tablosunda hiçbir hücrenin boş kalmamasını sağlar."""
    if deger is None or (isinstance(deger, str) and not deger.strip()):
        return "-"
    return deger


def main():
    json_path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("sonuclar.json")
    xlsx_path = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("fisler.xlsx")

    sonuclar = json.loads(json_path.read_text(encoding="utf-8"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Fişler"

    ws.append([baslik for _, baslik in SUTUNLAR])
    for hucre in ws[1]:
        hucre.font = Font(bold=True)

    supheli_sayisi = 0
    for veri in sonuclar:
        veri = kaydi_dogrula(fis_normalize_et(veri))
        satir = [excel_degeri(veri.get(anahtar)) for anahtar, _ in SUTUNLAR]
        ws.append(satir)
        if veri.get("_kontrol_gerekli"):
            supheli_sayisi += 1
            for hucre in ws[ws.max_row]:
                hucre.fill = SARI

    for col in ws.columns:
        uzunluk = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(uzunluk + 2, 40)

    wb.save(xlsx_path)
    print(f"{len(sonuclar)} satır '{xlsx_path}' dosyasına yazıldı.")
    print(f"{supheli_sayisi} satır kontrol gerektiriyor (sarı işaretli).")


if __name__ == "__main__":
    main()
