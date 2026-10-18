"""Excel aktarımının doğruluğunu etiketli (onaylı) gerçek fişlerle ölçer.

İki ölçüm:
- gercek: daha önce aktarılmış onaylı fişlerin defterdeki satırı
  (_excel_dosya/_excel_satir) okunur ve onaylı değerle karşılaştırılır.
  Fark, ya aktarım hatası ya da defterde sonradan yapılmış elle düzeltmedir.
- simule: tüm onaylı fişler bir defter kopyasına (gerçek defter verilirse
  onun kopyasına, yoksa boş deftere) uygulamanın aktarımıyla yazılır ve
  hücreler geri okunur. Gerçek defter değişmez.

Beklenen hücre değeri excel_aktar'dan bağımsız olarak burada hesaplanır.

    python -m tools.excel_dogruluk --defter "data/output/<defter>.xlsx" --sayfa "2026 MAZOT" \
        --cikti benchmark_excel.json
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import tempfile
from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook

from mazot import excel_aktar
from mazot.yollar import VERI_DIZINI

SUTUNLAR = {"tarih": 1, "fis_no": 2, "firma": 3, "plaka": 4, "ad_soyad": 5, "km": 6, "litre": 8}


def beklenen(alan: str, deger):
    """Onaylı metin değerin Excel'de görülmesi gereken hali."""
    if deger is None or str(deger).strip() in ("", "-"):
        return None
    metin = str(deger).strip()
    if alan == "tarih":
        m = re.fullmatch(r"(\d{1,2})[./](\d{1,2})[./](\d{4})", metin)
        if m:
            try:
                return datetime(int(m[3]), int(m[2]), int(m[1]))
            except ValueError:
                return metin
        return metin
    if alan in ("fis_no", "km"):
        sade = metin.replace(".", "").replace(" ", "")
        return int(sade) if sade.isdigit() else metin
    if alan == "litre":
        sade = metin.replace(" ", "")
        if "," in sade and "." in sade:
            ondalik = "," if sade.rfind(",") > sade.rfind(".") else "."
            sade = sade.replace("." if ondalik == "," else ",", "").replace(ondalik, ".")
        else:
            sade = sade.replace(",", ".")
        try:
            sayi = Decimal(sade)
        except InvalidOperation:
            return metin
        return int(sayi) if sayi == sayi.to_integral_value() else float(sayi)
    return metin


def esit(hucre, beklenen_deger) -> bool:
    if isinstance(hucre, (int, float)) and isinstance(beklenen_deger, (int, float)):
        return abs(hucre - beklenen_deger) < 1e-9
    if isinstance(hucre, datetime) and isinstance(beklenen_deger, datetime):
        return hucre.date() == beklenen_deger.date()
    return hucre == beklenen_deger


def onayli_kayitlar(kok: Path) -> list[dict]:
    kayitlar = []
    for yol in [kok / "sonuclar.json", *sorted((kok / "arsiv").glob("*/sonuclar.json"))]:
        if not yol.is_file():
            continue
        for kayit in json.loads(yol.read_text(encoding="utf-8")):
            if kayit.get("_onaylandi") and not kayit.get("_silindi") and not kayit.get("HATA"):
                kayit["_parti"] = yol.parent.name if yol.parent != kok else "aktif"
                kayitlar.append(kayit)
    return kayitlar


def _karsilastir(ornekler: list[tuple[dict, dict]]) -> dict:
    """ornekler: (kayıt, {alan: hücre değeri})."""
    dogru = Counter()
    farklar = []
    tam = 0
    for kayit, hucreler in ornekler:
        hepsi = True
        for alan in SUTUNLAR:
            b = beklenen(alan, kayit.get(alan))
            if esit(hucreler.get(alan), b):
                dogru[alan] += 1
            else:
                hepsi = False
                farklar.append({"dosya": kayit.get("_dosya"), "parti": kayit.get("_parti"),
                                "alan": alan, "onayli": kayit.get(alan),
                                "beklenen": str(b), "hucre": str(hucreler.get(alan))})
        tam += hepsi
    n = len(ornekler)
    return {
        "ornek": n,
        "tam_dogru_satir": round(tam / n * 100, 2) if n else None,
        "alan_bazinda": {a: round(dogru[a] / n * 100, 2) for a in SUTUNLAR} if n else {},
        "fark_sayisi": len(farklar),
        "fark_alan_dagilimi": dict(Counter(f["alan"] for f in farklar)),
        "farklar": farklar[:200],
    }


def gercek_olcum(kayitlar: list[dict], kok: Path) -> dict:
    defterler: dict[str, object] = {}
    ornekler, bulunamayan = [], Counter()
    for kayit in kayitlar:
        if not (kayit.get("_excel_aktarildi") and kayit.get("_excel_satir")):
            continue
        yol = Path(kayit.get("_excel_dosya") or "")
        if not yol.is_file():
            # Defter başka makinede/klasörde aktarılmış olabilir: adıyla veri klasöründe ara.
            adaylar = list((kok / "output").glob(yol.name)) if yol.name else []
            if not adaylar:
                bulunamayan[yol.name or "?"] += 1
                continue
            yol = adaylar[0]
        anahtar = str(yol.resolve())
        if anahtar not in defterler:
            defterler[anahtar] = load_workbook(yol, data_only=True)
        wb = defterler[anahtar]
        sayfa = kayit.get("_excel_sayfa")
        if sayfa not in wb.sheetnames:
            bulunamayan[f"{yol.name}/{sayfa}"] += 1
            continue
        ws = wb[sayfa]
        satir = int(kayit["_excel_satir"])
        ornekler.append((kayit, {a: ws.cell(row=satir, column=s).value for a, s in SUTUNLAR.items()}))
    for wb in defterler.values():
        wb.close()
    sonuc = _karsilastir(ornekler)
    sonuc["defteri_bulunamayan"] = dict(bulunamayan)
    return sonuc


def simule_olcum(kayitlar: list[dict], defter: Path | None, sayfa: str) -> dict:
    with tempfile.TemporaryDirectory() as gecici:
        kok = Path(gecici)
        hedef = kok / "output" / "kopya.xlsx"
        hedef.parent.mkdir()
        if defter:
            shutil.copy2(defter, hedef)
        else:
            wb = Workbook()
            wb.active.title = sayfa
            wb.active.append(["TARİH", "FATURA NO", "FİRMA", "PLAKA", "AD", "KM", None, "LİTRE"])
            wb.save(hedef)
        veri = kok / "sonuclar.json"
        temiz = [{k: v for k, v in kayit.items() if not k.startswith("_excel")}
                 for kayit in kayitlar]
        veri.write_text(json.dumps(temiz, ensure_ascii=False), encoding="utf-8")
        # Yedek ve "son kullanılan defter" ayarı da geçici klasöre yazılsın:
        # ölçüm kullanıcının verisinde iz bırakmamalı.
        eski = (excel_aktar.ROOT, excel_aktar.ARAMA_KLASORLERI, excel_aktar.AYAR_FILE)
        excel_aktar.ROOT, excel_aktar.ARAMA_KLASORLERI = kok, [kok / "output"]
        excel_aktar.AYAR_FILE = kok / "output" / ".excel_ayar.json"
        try:
            aktarim = excel_aktar.aktar(hedef, sayfa, tumu=True, data_file=veri)
        finally:
            excel_aktar.ROOT, excel_aktar.ARAMA_KLASORLERI, excel_aktar.AYAR_FILE = eski
        if "hata" in aktarim:
            return {"hata": aktarim["hata"]}
        yazilan = {k.get("_yukleme_id") or k.get("_dosya"): k
                   for k in json.loads(veri.read_text(encoding="utf-8"))}
        ws = load_workbook(hedef)[sayfa]
        ornekler = []
        for kayit in kayitlar:
            sonuc = yazilan.get(kayit.get("_yukleme_id") or kayit.get("_dosya"))
            if not sonuc or not sonuc.get("_excel_satir"):
                ornekler.append((kayit, {}))
                continue
            satir = sonuc["_excel_satir"]
            ornekler.append((kayit, {a: ws.cell(row=satir, column=s).value
                                     for a, s in SUTUNLAR.items()}))
        sonuc = _karsilastir(ornekler)
        sonuc["aktarilan"] = aktarim["aktarilan"]
        return sonuc


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--kok", type=Path, default=VERI_DIZINI)
    parser.add_argument("--defter", type=Path, help="simülasyon için kopyalanacak defter")
    parser.add_argument("--sayfa", default="2026 MAZOT")
    parser.add_argument("--cikti", type=Path)
    parser.add_argument("--esik", type=float, default=100.0,
                        help="simülasyonda tam doğru satır yüzdesi bunun altındaysa çıkış kodu 1")
    args = parser.parse_args()
    kayitlar = onayli_kayitlar(args.kok)
    rapor = {"etiketli_fis": len(kayitlar),
             "gercek": gercek_olcum(kayitlar, args.kok),
             "simule": simule_olcum(kayitlar, args.defter, args.sayfa)}
    ozet = {k: ({kk: vv for kk, vv in v.items() if kk != "farklar"} if isinstance(v, dict) else v)
            for k, v in rapor.items()}
    print(json.dumps(ozet, ensure_ascii=False, indent=2, default=str))
    if args.cikti:
        args.cikti.write_text(json.dumps(rapor, ensure_ascii=False, indent=2, default=str),
                              encoding="utf-8")
    basari = rapor["simule"].get("tam_dogru_satir")
    raise SystemExit(0 if basari is not None and basari >= args.esik else 1)


if __name__ == "__main__":
    main()
