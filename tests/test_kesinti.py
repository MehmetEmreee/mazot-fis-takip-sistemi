"""Kesinti testleri: işlem her dosya sistemi adımında öldürülür, toparlanmalı.

Her senaryo için gerçek bir alt süreç N'inci dosya işleminden (os.replace,
shutil.move, mkdir, yazma, silme, yeniden adlandırma) hemen sonra
os._exit(73) ile ölür; N=1'den işlem kendiliğinden bitene kadar denenir.
Her kesintiden sonra uygulamanın başlangıç kurtarması bu süreçte çalıştırılır
ve veri kaybı/çelişki olmadığı, işin güvenle tekrarlanabildiği doğrulanır.
Elektrik kesintisi (fsync/önbellek) testi değildir; süreç kesintisidir.
"""

from __future__ import annotations

import io
import json
import shutil
import subprocess
import sys
import tempfile
import textwrap
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from mazot import arayuz, excel_aktar, excel_ogrenme, parti_yonetimi

DEPO = Path(__file__).resolve().parent.parent
SAYFA = "2026 MAZOT"
UST_SINIR = 200  # sonsuz döngü koruması

KESICI = """
import os, pathlib, shutil, sys, json
N = int(sys.argv[1])
_sayac = [0]
def _nokta():
    _sayac[0] += 1
    if _sayac[0] == N:
        os._exit(73)
def _sar(sahip, ad):
    gercek = getattr(sahip, ad)
    def sarili(*a, **k):
        sonuc = gercek(*a, **k)
        _nokta()
        return sonuc
    setattr(sahip, ad, sarili)
for _sahip, _ad in ((os, "replace"), (os, "rename"), (shutil, "move"),
                    (pathlib.Path, "mkdir"), (pathlib.Path, "write_bytes"),
                    (pathlib.Path, "write_text"), (pathlib.Path, "unlink"),
                    (pathlib.Path, "rename"), (pathlib.Path, "rmdir")):
    _sar(_sahip, _ad)
"""


def jpeg(etiket: int) -> bytes:
    return b"\xff\xd8\xff\xe0JFIF" + etiket.to_bytes(4, "big") * 16 + b"\xff\xd9"


def sahte_ocr(yol) -> dict:
    return {"_dosya": Path(yol).name, "fis_no": "100001", "tarih": "05.05.2026",
            "firma": "F", "plaka": "34 ABC 123", "ad_soyad": "Ali Veli",
            "km": "12000", "litre": "50"}


class KesintiTemel(unittest.TestCase):
    senaryo_kodu = ""

    def setUp(self) -> None:
        self._gecici = tempfile.TemporaryDirectory()
        self.taban = Path(self._gecici.name).resolve()
        self._eski_arayuz = {a: getattr(arayuz, a) for a in ("ROOT", "DATA_FILE", "FISLER_DIR")}
        self._eski_excel = {a: getattr(excel_aktar, a) for a in (
            "ROOT", "DATA_FILE", "AYAR_FILE", "ARAMA_KLASORLERI")}
        self._eski_parti = (parti_yonetimi.ROOT, parti_yonetimi.DATA_FILE,
                            parti_yonetimi.FISLER_DIR)
        self._eski_hash = dict(arayuz.IMAGE_HASHES)

    def tearDown(self) -> None:
        for a, d in self._eski_arayuz.items():
            setattr(arayuz, a, d)
        for a, d in self._eski_excel.items():
            setattr(excel_aktar, a, d)
        parti_yonetimi.yollari_ayarla(*self._eski_parti)
        arayuz.IMAGE_HASHES.clear()
        arayuz.IMAGE_HASHES.update(self._eski_hash)
        self._gecici.cleanup()

    def baglan(self, kok: Path) -> None:
        arayuz.ROOT, arayuz.DATA_FILE, arayuz.FISLER_DIR = kok, kok / "sonuclar.json", kok / "fisler"
        excel_aktar.ROOT, excel_aktar.DATA_FILE = kok, kok / "sonuclar.json"
        excel_aktar.AYAR_FILE = kok / "output" / ".excel_ayar.json"
        excel_aktar.ARAMA_KLASORLERI = [kok / "output", kok]
        parti_yonetimi.yollari_ayarla(kok, kok / "sonuclar.json", kok / "fisler")

    def hazirla(self, kok: Path) -> None:
        raise NotImplementedError

    def baslangic_kurtarmasi(self) -> list:
        """arayuz.main()'in açılışta yaptığı uzlaştırma adımları."""
        with redirect_stdout(io.StringIO()):
            parti_yonetimi.yarim_parti_islemini_kurtar()
            arayuz.hash_indeksini_hazirla()
            return arayuz.yarim_ocr_kayitlarini_kurtar()

    def kes(self, kok: Path, n: int) -> int:
        kod = KESICI + textwrap.dedent(f"""
            from pathlib import Path
            from mazot import arayuz, excel_aktar, excel_ogrenme, parti_yonetimi, alan_yeniden_okuma
            kok = Path({str(kok)!r})
            arayuz.ROOT, arayuz.DATA_FILE, arayuz.FISLER_DIR = kok, kok / "sonuclar.json", kok / "fisler"
            excel_aktar.ROOT, excel_aktar.DATA_FILE = kok, kok / "sonuclar.json"
            excel_aktar.AYAR_FILE = kok / "output" / ".excel_ayar.json"
            excel_aktar.ARAMA_KLASORLERI = [kok / "output", kok]
            parti_yonetimi.yollari_ayarla(kok, kok / "sonuclar.json", kok / "fisler")
            arayuz.hash_indeksini_hazirla()
        """) + textwrap.dedent(self.senaryo_kodu)
        cikis = subprocess.run([sys.executable, "-c", kod, str(n)], cwd=DEPO,
                               capture_output=True, text=True, timeout=60)
        self.assertIn(cikis.returncode, (0, 73), cikis.stderr)
        return cikis.returncode

    def her_adimda_kes(self) -> int:
        for n in range(1, UST_SINIR):
            kok = self.taban / f"n{n:03d}"
            kok.mkdir()
            self.hazirla(kok)
            kod = self.kes(kok, n)
            self.baglan(kok)
            with self.subTest(kesinti_adimi=n, tamamlandi=kod == 0):
                self.dogrula(kok, kod == 0)
            if kod == 0:
                return n
        self.fail("işlem üst sınıra kadar bitmedi")

    def dogrula(self, kok: Path, tamamlandi: bool) -> None:
        raise NotImplementedError

    def kayitlar(self, yol: Path) -> list[dict]:
        return json.loads(yol.read_text(encoding="utf-8")) if yol.exists() else []


def onayli(ad: str, fis_no: str) -> dict:
    return {"_dosya": ad, "_yukleme_id": f"id-{ad}", "_onaylandi": True,
            "_dosya_hash": "h-" + ad, "fis_no": fis_no, "tarih": "05.05.2026",
            "firma": "F", "plaka": "34 ABC 123", "ad_soyad": "Ali Veli",
            "km": "12000", "litre": "50"}


class PartiKapatmaKesintisi(KesintiTemel):
    senaryo_kodu = 'parti_yonetimi.yeni_parti("Yeni Parti")'

    def hazirla(self, kok: Path) -> None:
        (kok / "fisler" / "F").mkdir(parents=True)
        (kok / "fisler" / "F" / "001.jpg").write_bytes(jpeg(1))
        (kok / "fisler" / "002.jpg").write_bytes(jpeg(2))
        (kok / "sonuclar.json").write_text(json.dumps(
            [onayli("001.jpg", "100001"), onayli("002.jpg", "100002")]), encoding="utf-8")
        (kok / "aktif_parti.json").write_text(
            json.dumps({"id": "aktif", "ad": "Eski Parti"}), encoding="utf-8")

    def dogrula(self, kok: Path, tamamlandi: bool) -> None:
        self.baslangic_kurtarmasi()
        self.baslangic_kurtarmasi()  # kurtarma tekrar çalıştırılabilir olmalı
        self.assertFalse(parti_yonetimi.ISLEM_GUNLUGU.exists())
        bolumler = [(kok / "sonuclar.json", kok / "fisler")] + [
            (a / "sonuclar.json", a / "fisler") for a in sorted((kok / "arsiv").glob("*"))
            if (kok / "arsiv").exists()]
        gorulen = {}
        for data, fisler in bolumler:
            self.assertTrue(data.exists(), f"{data} yok")
            for kayit in self.kayitlar(data):
                self.assertNotIn(kayit["_dosya"], gorulen, "kayıt iki partide")
                self.assertIsNotNone(arayuz.find_image(kayit["_dosya"], fisler),
                                     f"{kayit['_dosya']} görseli kendi partisinde yok")
                gorulen[kayit["_dosya"]] = data
        self.assertEqual(set(gorulen), {"001.jpg", "002.jpg"})
        # Partiye bağlı olmayan sahipsiz görsel kalmamalı.
        for gorsel in kok.rglob("*.jpg"):
            self.assertIn(gorsel.name, gorulen, f"sahipsiz görsel: {gorsel}")
        # Arşivler doğrulanabilir; işlem güvenle (yeniden) yapılabilir.
        for data, _ in bolumler[1:]:
            self.assertTrue(parti_yonetimi.arsiv_manifesti_dogrula(data.parent)["saglam"])
        if self.kayitlar(kok / "sonuclar.json"):
            parti_yonetimi.yeni_parti("Yeni Parti")
        self.assertEqual(parti_yonetimi.aktif_parti()["ad"], "Yeni Parti")

    def test_her_adimda_kesinti(self):
        self.assertGreater(self.her_adimda_kes(), 5)


class YuklemeKesintisi(KesintiTemel):
    senaryo_kodu = f'arayuz.fotografi_kabul_et({jpeg(7)!r}, ".jpg")'

    def hazirla(self, kok: Path) -> None:
        (kok / "fisler").mkdir()

    def dogrula(self, kok: Path, tamamlandi: bool) -> None:
        isler = self.baslangic_kurtarmasi()
        kayitlar = self.kayitlar(kok / "sonuclar.json")
        gorseller = list((kok / "fisler").rglob("*.jpg"))
        # Ya hiç iz yok ya da kayıt + görsel birlikte, kuyruğa alınmış.
        self.assertEqual(len(kayitlar), len(gorseller))
        self.assertLessEqual(len(kayitlar), 1)
        if kayitlar:
            self.assertEqual(kayitlar[0]["_dosya"], gorseller[0].name)
            self.assertEqual([p.name for p, _ in isler], [gorseller[0].name])
        if tamamlandi:
            self.assertEqual(len(kayitlar), 1)
        # Yeniden yükleme ya kabul edilir ya da var olan kayda işaret eder.
        with redirect_stdout(io.StringIO()):
            sonuc = arayuz.fotografi_kabul_et(jpeg(7), ".jpg")
        if kayitlar:
            self.assertEqual(sonuc.get("existing_filename"), kayitlar[0]["_dosya"])
        else:
            self.assertTrue(sonuc["ok"])

    def test_her_adimda_kesinti(self):
        self.assertGreater(self.her_adimda_kes(), 1)


class OcrIsciKesintisi(KesintiTemel):
    senaryo_kodu = """
        arayuz.resmi_oku = lambda p: {"_dosya": p.name, "fis_no": "100001",
            "tarih": "05.05.2026", "firma": "F", "plaka": "34 ABC 123",
            "ad_soyad": "Ali Veli", "km": "12000", "litre": "50"}
        excel_ogrenme.guvenli_otomatik_duzelt = lambda r: None
        alan_yeniden_okuma.yeniden_oku_ve_duzelt = lambda *a: None
        for is_ in arayuz.yarim_ocr_kayitlarini_kurtar():
            arayuz.fisi_isle(*is_)
    """

    def hazirla(self, kok: Path) -> None:
        self.baglan(kok)
        arayuz.IMAGE_HASHES.clear()
        with redirect_stdout(io.StringIO()):
            arayuz.fotografi_kabul_et(jpeg(3), ".jpg")

    def dogrula(self, kok: Path, tamamlandi: bool) -> None:
        isler = self.baslangic_kurtarmasi()
        with mock.patch.object(arayuz, "resmi_oku", side_effect=sahte_ocr), \
                mock.patch.object(arayuz.alan_yeniden_okuma, "yeniden_oku_ve_duzelt"), \
                mock.patch.object(excel_ogrenme, "guvenli_otomatik_duzelt"):
            for is_ in isler:
                arayuz.fisi_isle(*is_)
        kayitlar = self.kayitlar(kok / "sonuclar.json")
        self.assertEqual(len(kayitlar), 1)
        self.assertFalse(kayitlar[0]["_isleniyor"])
        self.assertNotIn("HATA", kayitlar[0])
        self.assertEqual(kayitlar[0]["_dosya"], "100001.jpg")
        self.assertEqual([p.name for p in (kok / "fisler").rglob("*.jpg")], ["100001.jpg"])

    def test_her_adimda_kesinti(self):
        self.assertGreater(self.her_adimda_kes(), 1)


class ExcelAktarimKesintisi(KesintiTemel):
    senaryo_kodu = 'excel_aktar.aktar(kok / "output" / "defter.xlsx", "2026 MAZOT")'

    def hazirla(self, kok: Path) -> None:
        (kok / "output").mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = SAYFA
        ws.append(["TARİH", "FATURA NO", "FİRMA", "PLAKA", "AD", "KM", None, "LİTRE"])
        ws.append(["01.05.2026", 100, "X", "34 X 1", "Y", 10, None, 5])
        wb.save(kok / "output" / "defter.xlsx")
        (kok / "sonuclar.json").write_text(json.dumps(
            [onayli("001.jpg", "100001"), onayli("002.jpg", "100002")]), encoding="utf-8")

    def dogrula(self, kok: Path, tamamlandi: bool) -> None:
        defter = kok / "output" / "defter.xlsx"
        sonuc = excel_aktar.aktar(defter, SAYFA)  # kullanıcının tekrar denemesi
        self.assertNotIn("hata", sonuc)
        ws = load_workbook(defter)[SAYFA]
        fisler = [ws.cell(row=r, column=2).value for r in range(3, ws.max_row + 1)]
        self.assertEqual(sorted(f for f in fisler if f), [100001, 100002])
        for kayit in self.kayitlar(kok / "sonuclar.json"):
            self.assertTrue(kayit["_excel_aktarildi"])
            self.assertNotIn("_excel_bekleyen", kayit)
            self.assertEqual(ws.cell(row=kayit["_excel_satir"], column=2).value,
                             int(kayit["fis_no"]))
        self.assertEqual([p.name for p in (kok / "output").glob(".*.xlsx")], [])

    def test_her_adimda_kesinti(self):
        self.assertGreater(self.her_adimda_kes(), 3)


if __name__ == "__main__":
    unittest.main()
