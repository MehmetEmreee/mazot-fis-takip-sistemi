"""f46cdc9 hata araştırmasındaki H01-H09 bulgularının kabul testleri.

Ollama gerekmez: OCR ve öğrenme adımları sahteyle değiştirilir. H02 ve H07
gerçek bir alt süreci kritik adımda os._exit(73) ile sonlandırır, ardından
başlangıç kurtarmasını bu süreçte çalıştırır.
Çalıştırma: .venv/bin/python -m unittest tests.test_hata_duzeltmeleri
"""

from __future__ import annotations

import hashlib
import io
import json
import queue
import subprocess
import sys
import tempfile
import textwrap
import threading
import time
import unittest
from contextlib import redirect_stdout
from http import HTTPStatus
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

from mazot import arayuz, denetim, erisim
from mazot import excel_aktar
from mazot import excel_ogrenme
from mazot import parti_yonetimi

DEPO = Path(__file__).resolve().parent.parent
SAYFA = "2026 MAZOT"


def jpeg(etiket: int) -> bytes:
    return b"\xff\xd8\xff\xe0JFIF" + etiket.to_bytes(4, "big") * 16 + b"\xff\xd9"


def onayli(ad: str, fis_no: str, **ek) -> dict:
    return {
        "_dosya": ad, "_yukleme_id": f"id-{ad}", "_onaylandi": True,
        "fis_no": fis_no, "tarih": "05.05.2026", "firma": "Selim İnşaat",
        "plaka": "34 ABC 123", "ad_soyad": "Ali Veli", "km": "12000",
        "litre": "50", **ek,
    }


def defter_olustur(yol: Path, a1: str = "TARİH") -> Path:
    yol.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = SAYFA
    ws.append([a1, "FATURA NO", "FİRMA", "PLAKA", "AD", "KM", None, "LİTRE"])
    ws.append(["01.05.2026", 100, "X", "34 X 1", "Y", 10, None, 5])
    wb.save(yol)
    return yol


class SahteHandler(arayuz.ReceiptHandler):
    """Soket açmadan gerçek handler metotlarını çağırır."""

    def __init__(self, path: str = "/", govde: dict | None = None) -> None:
        veri = json.dumps(govde).encode("utf-8") if govde is not None else b""
        self.path = path
        self.client_address = ("127.0.0.1", 0)
        self.headers = {"Content-Length": str(len(veri))}
        self.rfile = io.BytesIO(veri)
        self.yanitlar: list[tuple[int, dict]] = []

    def send_json(self, payload, status=HTTPStatus.OK) -> None:
        self.yanitlar.append((int(status), payload))


class Temel(unittest.TestCase):
    def setUp(self) -> None:
        self._gecici = tempfile.TemporaryDirectory()
        self.kok = Path(self._gecici.name).resolve()
        self.data = self.kok / "sonuclar.json"
        self.fisler = self.kok / "fisler"
        self.fisler.mkdir()
        self._eski_arayuz = {a: getattr(arayuz, a) for a in (
            "ROOT", "DATA_FILE", "FISLER_DIR", "OCR_QUEUE")}
        self._eski_excel = {a: getattr(excel_aktar, a) for a in (
            "ROOT", "DATA_FILE", "AYAR_FILE", "ARAMA_KLASORLERI")}
        self._eski_parti = (parti_yonetimi.ROOT, parti_yonetimi.DATA_FILE,
                            parti_yonetimi.FISLER_DIR)
        self._eski_hash = dict(arayuz.IMAGE_HASHES)
        arayuz.ROOT, arayuz.DATA_FILE, arayuz.FISLER_DIR = self.kok, self.data, self.fisler
        arayuz.OCR_QUEUE = queue.Queue()
        arayuz.IMAGE_HASHES.clear()
        excel_aktar.ROOT = self.kok
        excel_aktar.DATA_FILE = self.data
        excel_aktar.AYAR_FILE = self.kok / "output" / ".excel_ayar.json"
        excel_aktar.ARAMA_KLASORLERI = [self.kok / "output", self.kok]
        parti_yonetimi.yollari_ayarla(self.kok, self.data, self.fisler)
        self._eski_denetim, self._eski_erisim = denetim.DENETIM_DIR, erisim.KULLANICI_FILE
        denetim.yolu_ayarla(self.kok)
        erisim.yolu_ayarla(self.kok)

    def tearDown(self) -> None:
        denetim.DENETIM_DIR, erisim.KULLANICI_FILE = self._eski_denetim, self._eski_erisim
        for a, d in self._eski_arayuz.items():
            setattr(arayuz, a, d)
        for a, d in self._eski_excel.items():
            setattr(excel_aktar, a, d)
        parti_yonetimi.yollari_ayarla(*self._eski_parti)
        arayuz.IMAGE_HASHES.clear()
        arayuz.IMAGE_HASHES.update(self._eski_hash)
        self._gecici.cleanup()

    def yaz(self, yol: Path, kayitlar: list[dict]) -> None:
        yol.parent.mkdir(parents=True, exist_ok=True)
        yol.write_text(json.dumps(kayitlar, ensure_ascii=False), encoding="utf-8")

    def oku(self, yol: Path | None = None) -> list[dict]:
        return json.loads((yol or self.data).read_text(encoding="utf-8"))

    def alt_surec(self, kod: str) -> subprocess.CompletedProcess:
        return subprocess.run(
            [sys.executable, "-c", textwrap.dedent(kod)], cwd=DEPO,
            capture_output=True, text=True, timeout=60,
        )


class H09LitreOndalikTesti(Temel):
    def test_alan_bazinda_sayi_donusumu(self):
        self.assertEqual(excel_aktar.sayiya_cevir("12.5", "litre"), 12.5)
        self.assertEqual(excel_aktar.sayiya_cevir("12,5", "litre"), 12.5)
        self.assertEqual(excel_aktar.sayiya_cevir("1.234,5", "litre"), 1234.5)
        self.assertEqual(excel_aktar.sayiya_cevir("1,234.5", "litre"), 1234.5)
        self.assertEqual(excel_aktar.sayiya_cevir("50", "litre"), 50)
        # km ve fiş no'da nokta binlik ayracı olarak kalır.
        self.assertEqual(excel_aktar.sayiya_cevir("12.500", "km"), 12500)
        self.assertEqual(excel_aktar.sayiya_cevir("123456", "fis_no"), 123456)
        # Gerçek veride bulundu: boşlukla ayrılmış km metin olarak yazılıyordu.
        self.assertEqual(excel_aktar.sayiya_cevir("300 638", "km"), 300638)
        self.assertEqual(excel_aktar.sayiya_cevir("300\u00a0638", "km"), 300638)

    def test_gercek_defter_hucresi(self):
        defter = defter_olustur(self.kok / "output" / "defter.xlsx")
        self.yaz(self.data, [onayli("a.jpg", "100001", litre="12.5", km="12.500")])
        sonuc = excel_aktar.aktar(defter, SAYFA, data_file=self.data)
        self.assertEqual(sonuc["aktarilan"], 1, sonuc)
        ws = load_workbook(defter)[SAYFA]
        self.assertEqual(ws.cell(row=sonuc["ilk_satir"], column=8).value, 12.5)
        self.assertEqual(ws.cell(row=sonuc["ilk_satir"], column=6).value, 12500)


class H08YedekCakismasiTesti(Temel):
    def test_ayni_adli_farkli_defterler_ayri_yedeklenir(self):
        a = defter_olustur(self.kok / "output" / "defter.xlsx", a1="BIRINCI")
        b = defter_olustur(self.kok / "baska" / "defter.xlsx", a1="IKINCI")
        ya = excel_aktar.aktarim_oncesi_yedekle(a)
        yb = excel_aktar.aktarim_oncesi_yedekle(b)
        self.assertNotEqual(ya, yb)
        self.assertEqual(load_workbook(ya)[SAYFA]["A1"].value, "BIRINCI")
        self.assertEqual(load_workbook(yb)[SAYFA]["A1"].value, "IKINCI")
        # Aynı hedef gün içinde tek yedek.
        self.assertEqual(excel_aktar.aktarim_oncesi_yedekle(a), ya)


class H01ParalelPartiOkumaTesti(Temel):
    def test_aktarim_sirasinda_arsiv_istegi_veri_yolunu_degistirmez(self):
        arsiv = self.kok / "arsiv" / "parti_eski"
        self.yaz(arsiv / "sonuclar.json", [onayli("eski.jpg", "100009")])
        arsiv_once = (arsiv / "sonuclar.json").read_bytes()
        self.yaz(self.data, [onayli("aktif.jpg", "100001")])
        defter = defter_olustur(self.kok / "output" / "defter.xlsx")

        gercek = excel_aktar.yarim_aktarimlari_uzlastir
        yanitlar = []

        def araya_gir(kayitlar):
            # Aktarım JSON'u okuduktan sonra arşiv partisi için Excel listesi istenir.
            handler = SahteHandler("/api/excel/files?parti=parti_eski")
            with mock.patch.object(excel_ogrenme, "durum", return_value={}):
                handler.do_GET()
            yanitlar.extend(handler.yanitlar)
            return gercek(kayitlar)

        with mock.patch.object(excel_aktar, "yarim_aktarimlari_uzlastir", araya_gir):
            handler = SahteHandler("/api/export-excel", {
                "dosya": str(defter), "sayfa": SAYFA, "parti": "aktif"})
            handler.do_POST()

        self.assertEqual(yanitlar[0][0], 200)
        self.assertEqual(yanitlar[0][1]["bekleyen"], 1)
        self.assertEqual(handler.yanitlar[0][0], 200, handler.yanitlar)
        self.assertEqual((arsiv / "sonuclar.json").read_bytes(), arsiv_once)
        self.assertTrue(self.oku()[0]["_excel_aktarildi"])
        self.assertIs(excel_aktar.DATA_FILE, self.data)


class H05IdempotentAktarimTesti(Temel):
    def setUp(self) -> None:
        super().setUp()
        self.defter = defter_olustur(self.kok / "output" / "defter.xlsx")
        self.yaz(self.data, [onayli("a.jpg", "100001")])

    def yarida_kes(self, **ek) -> dict:
        """Defter kaydedildikten sonra JSON tamamlama yazımı OSError verir."""
        gercek = excel_aktar.kaydet_json
        sayac = {"n": 0}

        def bozuk(yol, veri):
            if Path(yol) == self.data:
                sayac["n"] += 1
                if sayac["n"] == 2:
                    raise OSError("sentetik disk hatası")
            gercek(yol, veri)

        with mock.patch.object(excel_aktar, "kaydet_json", bozuk):
            return excel_aktar.aktar(self.defter, SAYFA, data_file=self.data, **ek)

    def fis_satirlari(self, defter: Path | None = None) -> list[int]:
        ws = load_workbook(defter or self.defter)[SAYFA]
        return [r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=2).value == 100001]

    def test_yarim_kalan_aktarim_tekrar_denenince_cift_yazmaz(self):
        self.assertIn("hata", self.yarida_kes())
        self.assertEqual(len(self.fis_satirlari()), 1)
        self.assertIn("_excel_bekleyen", self.oku()[0])

        sonuc = excel_aktar.aktar(self.defter, SAYFA, data_file=self.data)
        self.assertEqual((sonuc["aktarilan"], sonuc.get("tamamlanan")), (0, 1), sonuc)
        self.assertEqual(len(self.fis_satirlari()), 1)
        kayit = self.oku()[0]
        self.assertTrue(kayit["_excel_aktarildi"])
        self.assertNotIn("_excel_bekleyen", kayit)
        self.assertEqual(kayit["_excel_satir"], self.fis_satirlari()[0])
        # Defter sayfaları/görünümü: kimlik sayfası gizli, kullanıcı listesinde yok.
        wb = load_workbook(self.defter)
        self.assertEqual(wb[excel_aktar.KIMLIK_SAYFASI].sheet_state, "hidden")
        self.assertEqual(excel_aktar.sayfa_adlari(self.defter), [SAYFA])

    def test_tumu_ile_ayni_hedefe_yarim_tekrar_cift_yazmaz(self):
        self.yarida_kes(tumu=True)
        excel_aktar.aktar(self.defter, SAYFA, data_file=self.data, tumu=True)
        self.assertEqual(len(self.fis_satirlari()), 1)

    def test_kasitli_yeniden_aktarim_ve_farkli_hedef_korunur(self):
        self.yarida_kes()
        baska = defter_olustur(self.kok / "output" / "baska.xlsx")
        sonuc = excel_aktar.aktar(baska, SAYFA, data_file=self.data, tumu=True)
        self.assertEqual(sonuc["aktarilan"], 1)
        self.assertEqual(len(self.fis_satirlari(baska)), 1)
        # Tamamlanmış aktarım bilerek aynı deftere yeniden yazılabilir.
        sonuc = excel_aktar.aktar(self.defter, SAYFA, data_file=self.data, tumu=True)
        self.assertEqual(sonuc["aktarilan"], 1)
        self.assertEqual(len(self.fis_satirlari()), 2)

    def test_degistirilmis_bekleyen_satir_acik_hata_verir(self):
        self.yarida_kes()
        satir = self.oku()[0]["_excel_bekleyen"]["satir"]
        wb = load_workbook(self.defter)
        wb[SAYFA].cell(row=satir, column=8).value = 999
        wb.save(self.defter)
        sonuc = excel_aktar.aktar(self.defter, SAYFA, data_file=self.data)
        self.assertIn("değiştirilmiş", sonuc["hata"])
        self.assertIn("_excel_bekleyen", self.oku()[0])

    def test_baska_ait_kimlik_sayfasi_acik_hata_verir(self):
        wb = load_workbook(self.defter)
        wb.create_sheet(excel_aktar.KIMLIK_SAYFASI)["A1"] = "kullanıcı verisi"
        wb.save(self.defter)
        sonuc = excel_aktar.aktar(self.defter, SAYFA, data_file=self.data)
        self.assertIn("hata", sonuc)
        self.assertEqual(self.fis_satirlari(), [])
        self.assertNotIn("_excel_bekleyen", self.oku()[0])

    def test_defter_kaydi_basarisizsa_bekleme_temizlenir(self):
        with mock.patch.object(excel_aktar, "defteri_kaydet", side_effect=PermissionError):
            self.assertIn("hata", excel_aktar.aktar(self.defter, SAYFA, data_file=self.data))
        self.assertNotIn("_excel_bekleyen", self.oku()[0])
        self.assertEqual(excel_aktar.aktar(self.defter, SAYFA, data_file=self.data)["aktarilan"], 1)


class OcrTemel(Temel):
    def setUp(self) -> None:
        super().setUp()
        yamalar = [
            mock.patch.object(arayuz, "resmi_oku", side_effect=self.sahte_ocr),
            mock.patch.object(arayuz.alan_yeniden_okuma, "yeniden_oku_ve_duzelt"),
            mock.patch.object(excel_ogrenme, "guvenli_otomatik_duzelt"),
        ]
        self.ogrenme = yamalar[2].start()
        for yama in yamalar[:2]:
            yama.start()
        for yama in yamalar:
            self.addCleanup(yama.stop)
        self.ocr_bariyeri: threading.Event | None = None
        self.ocr_basladi = threading.Event()

    def sahte_ocr(self, image_path: Path) -> dict:
        self.ocr_basladi.set()
        if self.ocr_bariyeri:
            self.ocr_bariyeri.wait(5)
        numara = {1: "100001", 2: "100002"}.get(image_path.read_bytes()[11], "100003")
        return {"_dosya": image_path.name, "fis_no": numara, "firma": "OCR FİRMA",
                "tarih": "05.05.2026", "plaka": "34 ABC 123", "ad_soyad": "Ali Veli",
                "km": "12000", "litre": "50"}

    def isciyi_baslat(self) -> None:
        threading.Thread(target=arayuz.ocr_worker, args=(arayuz.OCR_QUEUE,), daemon=True).start()

    def kuyrugu_bekle(self, sure: float = 5) -> None:
        bitis = time.monotonic() + sure
        while arayuz.OCR_QUEUE.unfinished_tasks and time.monotonic() < bitis:
            time.sleep(0.02)
        self.assertEqual(arayuz.OCR_QUEUE.unfinished_tasks, 0, "kuyruk takıldı")


class H03IsciHataSiniriTesti(OcrTemel):
    def test_ogrenme_hatasi_isciyi_durdurmaz(self):
        self.ogrenme.side_effect = OSError("öğrenme dosyası kilitli")
        with redirect_stdout(io.StringIO()):
            arayuz.fotografi_kabul_et(jpeg(1), ".jpg")
            arayuz.fotografi_kabul_et(jpeg(2), ".jpg")
            self.isciyi_baslat()
            self.kuyrugu_bekle()
        kayitlar = self.oku()
        self.assertEqual([k["_isleniyor"] for k in kayitlar], [False, False])
        self.assertEqual({k["fis_no"] for k in kayitlar}, {"100001", "100002"})
        self.assertTrue(all(k.get("_uyarilar") for k in kayitlar))

    def test_beklenmeyen_hata_kaydi_isleniyorda_birakmaz(self):
        with redirect_stdout(io.StringIO()):
            arayuz.fotografi_kabul_et(jpeg(1), ".jpg")
            arayuz.fotografi_kabul_et(jpeg(2), ".jpg")
            gercek = arayuz.benzersiz_kayit_adi
            cagri = {"n": 0}

            def ilkinde_patla(*a, **k):
                cagri["n"] += 1
                if cagri["n"] == 1:
                    raise RuntimeError("sentetik")
                return gercek(*a, **k)

            with mock.patch.object(arayuz, "benzersiz_kayit_adi", ilkinde_patla):
                self.isciyi_baslat()
                self.kuyrugu_bekle()
        kayitlar = self.oku()
        self.assertFalse(any(k["_isleniyor"] for k in kayitlar))
        self.assertIn("sentetik", kayitlar[0]["HATA"])


class IsciDayaniklilikTesti(OcrTemel):
    def test_bozuk_kuyruk_ogesi_isciyi_oldurmez(self):
        with redirect_stdout(io.StringIO()):
            arayuz.OCR_QUEUE.put((None, "bozuk"))
            arayuz.OCR_QUEUE.put(("yok.jpg",))
            arayuz.fotografi_kabul_et(jpeg(1), ".jpg")
            self.isciyi_baslat()
            self.kuyrugu_bekle()
        self.assertFalse(self.oku()[0]["_isleniyor"])
        self.assertEqual(self.oku()[0]["fis_no"], "100001")


class H04IslenenKayitKorumaTesti(OcrTemel):
    def test_ocr_surerken_kaydet_ve_sil_reddedilir(self):
        self.ocr_bariyeri = threading.Event()
        with redirect_stdout(io.StringIO()):
            ad = arayuz.fotografi_kabul_et(jpeg(1), ".jpg")["filename"]
            self.isciyi_baslat()
            self.assertTrue(self.ocr_basladi.wait(5))
            kaydet = SahteHandler("/api/save", {
                "filename": ad, "fields": {"firma": "DOĞRU FİRMA"}, "approved": True})
            kaydet.do_POST()
            sil = SahteHandler("/api/delete", {"filename": ad})
            sil.do_POST()
            self.ocr_bariyeri.set()
            self.kuyrugu_bekle()
        self.assertEqual(kaydet.yanitlar[0][0], 409)
        self.assertEqual(sil.yanitlar[0][0], 409)
        kayit = self.oku()[0]
        self.assertFalse(kayit.get("_silindi"))
        self.assertEqual(kayit["firma"], "OCR FİRMA")
        # OCR bittikten sonra düzeltme kabul edilir.
        ad = kayit["_dosya"]
        with mock.patch.object(excel_ogrenme, "duzeltmelerden_ogren", return_value=0), \
                mock.patch.object(excel_ogrenme, "kayit_onerileri", return_value={}):
            kaydet = SahteHandler("/api/save", {
                "filename": ad, "fields": {"firma": "DOĞRU FİRMA"}, "approved": True})
            kaydet.do_POST()
        self.assertEqual(kaydet.yanitlar[0][0], 200, kaydet.yanitlar)
        self.assertEqual(self.oku()[0]["firma"], "DOĞRU FİRMA")
        self.assertTrue(self.oku()[0]["_onaylandi"])


class H06YuklemeVePartiKapatmaTesti(Temel):
    def test_yukleme_ile_parti_kapatma_ayni_kritik_bolumde(self):
        self.yaz(self.data, [onayli("eski.jpg", "100001")])
        (self.fisler / "eski.jpg").write_bytes(jpeg(9))
        gercek = arayuz.save_receipts
        kapatma = SahteHandler("/api/partiler/yeni", {"ad": "Yeni"})
        is_parcacigi: list[threading.Thread] = []
        beklerken_canli: list[bool] = []

        def kilitsiz_bosluk(receipts):
            # Yükleme kritik bölümündeyken parti kapatma isteği gelir.
            t = threading.Thread(target=kapatma.do_POST)
            t.start()
            t.join(0.3)
            is_parcacigi.append(t)
            beklerken_canli.append(t.is_alive())
            gercek(receipts)

        with mock.patch.object(arayuz, "save_receipts", side_effect=kilitsiz_bosluk):
            sonuc = arayuz.fotografi_kabul_et(jpeg(1), ".jpg")
        self.assertEqual(beklerken_canli, [True], "parti kapatma kilidi beklemeliydi")
        is_parcacigi[0].join(5)
        # Yeni yükleme işleniyor olduğu için parti kapatılmaz; ikisi aynı partide.
        self.assertEqual(kapatma.yanitlar[0][0], 400)
        kayit = [k for k in self.oku() if k["_dosya"] == sonuc["filename"]]
        self.assertEqual(len(kayit), 1)
        self.assertIsNotNone(arayuz.find_image(sonuc["filename"]))
        self.assertFalse((self.kok / "arsiv").exists())


class H02YenidenAdlandirmaKesintisiTesti(Temel):
    def test_yeniden_adlandirmadan_sonra_surec_olurse_kurtarilir(self):
        with redirect_stdout(io.StringIO()):
            sonuc = arayuz.fotografi_kabul_et(jpeg(1), ".jpg")
        ad = sonuc["filename"]
        kod = f"""
            import os
            from pathlib import Path
            from mazot import arayuz
            kok = Path({str(self.kok)!r})
            arayuz.DATA_FILE, arayuz.FISLER_DIR = kok / "sonuclar.json", kok / "fisler"
            arayuz.resmi_oku = lambda p: {{"_dosya": p.name, "fis_no": "100001",
                "tarih": "05.05.2026", "plaka": "34 ABC 123", "firma": "F",
                "ad_soyad": "Ali Veli", "km": "12000", "litre": "50"}}
            arayuz.excel_ogrenme.guvenli_otomatik_duzelt = lambda r: None
            arayuz.alan_yeniden_okuma.yeniden_oku_ve_duzelt = lambda *a: None
            arayuz.save_receipts = lambda r: os._exit(73)  # rename yapıldı, JSON yazılmadan öl
            p = arayuz.temp_dir() / {ad!r}
            arayuz.fisi_isle(p, {sonuc['receipt']['_dosya_hash']!r})
        """
        cikis = self.alt_surec(kod)
        self.assertEqual(cikis.returncode, 73, cikis.stderr)
        self.assertTrue((self.fisler / "temp" / "100001.jpg").is_file())
        self.assertTrue(self.oku()[0]["_isleniyor"])
        self.assertIsNone(arayuz.find_image(ad))

        # Eski sürümden kalan, görseli ve hash'i olmayan yarım kayıt da eklenir.
        kayitlar = self.oku() + [{"_dosya": "kayip.jpg", "_isleniyor": True}]
        self.yaz(self.data, kayitlar)
        arayuz.hash_indeksini_hazirla()
        isler = arayuz.yarim_ocr_kayitlarini_kurtar()
        self.assertEqual([p.name for p, _ in isler], ["100001.jpg"])
        kayitlar = self.oku()
        self.assertEqual(kayitlar[0]["_dosya"], "100001.jpg")
        self.assertTrue(kayitlar[0]["_isleniyor"])
        self.assertFalse(kayitlar[1]["_isleniyor"])
        self.assertIn("Görsel bulunamadı", kayitlar[1]["HATA"])
        # Tekrar kuyruğa alınan görsel yeniden _2 adını almaz.
        with mock.patch.object(arayuz, "resmi_oku", return_value={
                "fis_no": "100001", "tarih": "05.05.2026", "plaka": "34 ABC 123",
                "firma": "F", "ad_soyad": "Ali Veli", "km": "12000", "litre": "50"}), \
                mock.patch.object(arayuz.alan_yeniden_okuma, "yeniden_oku_ve_duzelt"), \
                mock.patch.object(excel_ogrenme, "guvenli_otomatik_duzelt"):
            arayuz.fisi_isle(*isler[0])
        self.assertEqual(self.oku()[0]["_dosya"], "100001.jpg")
        self.assertFalse(self.oku()[0]["_isleniyor"])


class H07PartiKapatmaKesintisiTesti(Temel):
    def hazirla(self) -> None:
        self.yaz(self.data, [onayli("001.jpg", "100001")])
        (self.fisler / "001.jpg").write_bytes(b"gorsel")
        (parti_yonetimi.AKTIF_META_FILE).write_text(
            json.dumps({"id": "aktif", "ad": "Eski Parti"}), encoding="utf-8")

    def kes(self, yama: str) -> None:
        kod = f"""
            import os
            from pathlib import Path
            from mazot import parti_yonetimi
            kok = Path({str(self.kok)!r})
            parti_yonetimi.yollari_ayarla(kok, kok / "sonuclar.json", kok / "fisler")
            {yama}
            parti_yonetimi.yeni_parti("Yeni Parti")
        """
        cikis = self.alt_surec(kod)
        self.assertEqual(cikis.returncode, 73, cikis.stderr)

    def arsivler(self) -> list[Path]:
        kok = self.kok / "arsiv"
        return [p for p in kok.iterdir()] if kok.exists() else []

    def test_veri_tasindiktan_sonra_kesinti_geri_alinir(self):
        self.hazirla()
        self.kes("parti_yonetimi.shutil.move = lambda *a: os._exit(73)")
        self.assertFalse(self.data.exists())
        self.assertEqual(parti_yonetimi.aktif_parti()["toplam"], 0)

        sonuc = parti_yonetimi.yarim_parti_islemini_kurtar()
        self.assertEqual(sonuc["durum"], "geri_alindi")
        self.assertEqual(self.oku()[0]["_dosya"], "001.jpg")
        self.assertTrue((self.fisler / "001.jpg").is_file())
        self.assertEqual(self.arsivler(), [])
        self.assertFalse(parti_yonetimi.ISLEM_GUNLUGU.exists())
        self.assertEqual(parti_yonetimi.aktif_parti()["ad"], "Eski Parti")
        # Tekrar kurtarma etkisizdir; işlem sonra normal tamamlanabilir.
        self.assertIsNone(parti_yonetimi.yarim_parti_islemini_kurtar())
        parti_yonetimi.yeni_parti("Yeni Parti")
        self.assertEqual(len(self.arsivler()), 1)

    def test_gorseller_tasindiktan_sonra_kesinti_geri_alinir(self):
        self.hazirla()
        self.kes("gercek = parti_yonetimi._json_yaz; "
                 "parti_yonetimi._json_yaz = lambda yol, veri: os._exit(73) "
                 "if yol.name == 'parti.json' else gercek(yol, veri)")
        self.assertFalse(self.fisler.exists())
        sonuc = parti_yonetimi.yarim_parti_islemini_kurtar()
        self.assertEqual(sonuc["durum"], "geri_alindi")
        self.assertTrue((self.fisler / "001.jpg").is_file())
        self.assertEqual(len(self.oku()), 1)

    def test_arsivlendikten_sonra_kesinti_ileri_tamamlanir(self):
        self.hazirla()
        self.kes("parti_yonetimi._yeni_aktifi_baslat = lambda yeni: os._exit(73)")
        sonuc = parti_yonetimi.yarim_parti_islemini_kurtar()
        self.assertEqual(sonuc["durum"], "tamamlandi")
        self.assertEqual(self.oku(), [])
        self.assertEqual(parti_yonetimi.aktif_parti()["ad"], "Yeni Parti")
        arsiv = self.arsivler()[0]
        self.assertTrue((arsiv / "fisler" / "001.jpg").is_file())
        self.assertTrue(parti_yonetimi.arsiv_manifesti_dogrula(arsiv)["saglam"])

    def test_celisen_dosyalarin_uzerine_yazilmaz(self):
        self.hazirla()
        self.kes("parti_yonetimi.shutil.move = lambda *a: os._exit(73)")
        # Kesintiden sonra (ör. elle) farklı içerikli aktif JSON oluşmuş.
        self.yaz(self.data, [{"_dosya": "baska.jpg"}])
        sonuc = parti_yonetimi.yarim_parti_islemini_kurtar()
        self.assertEqual(sonuc["durum"], "cakisma")
        self.assertEqual(self.oku()[0]["_dosya"], "baska.jpg")
        arsiv_json = self.arsivler()[0] / "sonuclar.json"
        self.assertEqual(json.loads(arsiv_json.read_text(encoding="utf-8"))[0]["_dosya"], "001.jpg")
        self.assertTrue(parti_yonetimi.ISLEM_GUNLUGU.exists())
        with self.assertRaises(parti_yonetimi.PartiIslemHatasi):
            parti_yonetimi.yeni_parti("Başka")

    def test_farkli_diskte_yarim_kopya_birlestirilir(self):
        # shutil.move farklı diskte kopyala+sil yapar: iki tarafta da dosya kalabilir.
        self.hazirla()
        hedef = self.kok / "arsiv" / "parti_yarim"
        (hedef / "fisler").mkdir(parents=True)
        (hedef / "fisler" / "001.jpg").write_bytes(b"gorsel")       # tam kopya
        (hedef / "fisler" / "002.jpg").write_bytes(b"sadece-arsiv")  # kaynaktan silinmiş
        parti_yonetimi._json_yaz(parti_yonetimi.ISLEM_GUNLUGU, {
            "hedef": str(hedef), "asama": "basladi", "yeni": {"id": "aktif", "ad": "Y"}})
        sonuc = parti_yonetimi.yarim_parti_islemini_kurtar()
        self.assertEqual(sonuc["durum"], "geri_alindi")
        self.assertEqual((self.fisler / "001.jpg").read_bytes(), b"gorsel")
        self.assertEqual((self.fisler / "002.jpg").read_bytes(), b"sadece-arsiv")
        self.assertFalse(hedef.exists())


if __name__ == "__main__":
    unittest.main()
