"""Erişim kontrolü (roller, oturum, telefon anahtarı) ve denetim kaydı testleri.

Gerçek handler metotları soket açmadan çağrılır.
"""

from __future__ import annotations

import io
import json
from contextlib import redirect_stdout
from http import HTTPStatus
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from mazot import arayuz, denetim, erisim, excel_aktar, excel_ogrenme
from tests.test_hata_duzeltmeleri import SAYFA, Temel, defter_olustur, jpeg, onayli

YEREL = "127.0.0.1"
AGDAN = "192.168.1.50"


class Istek(arayuz.ReceiptHandler):
    def __init__(self, yol: str, govde=None, istemci: str = YEREL,
                 cerez: str | None = None, ham: bytes | None = None,
                 basliklar: dict | None = None) -> None:
        veri = ham if ham is not None else (
            json.dumps(govde).encode("utf-8") if govde is not None else b"")
        self.path = yol
        self.client_address = (istemci, 0)
        self.headers = {"Content-Length": str(len(veri)), **(basliklar or {})}
        if cerez:
            self.headers["Cookie"] = cerez
        self.rfile = io.BytesIO(veri)
        self.yanitlar: list[tuple[int, dict]] = []
        self.cerezler: list[str] = []
        self.yonlendirmeler: list[str] = []

    def send_json(self, payload, status=HTTPStatus.OK, cerez=None) -> None:
        self.yanitlar.append((int(status), payload))
        if cerez:
            self.cerezler.append(cerez)

    def yonlendir(self, hedef, cerez=None) -> None:
        self.yonlendirmeler.append(hedef)
        if cerez:
            self.cerezler.append(cerez)

    @property
    def durum(self) -> int:
        return self.yanitlar[0][0]


def get(yol, **k) -> Istek:
    istek = Istek(yol, **k)
    with redirect_stdout(io.StringIO()):
        istek.do_GET()
    return istek


def post(yol, govde=None, **k) -> Istek:
    istek = Istek(yol, govde, **k)
    with redirect_stdout(io.StringIO()), \
            mock.patch.object(excel_ogrenme, "duzeltmelerden_ogren", return_value=0), \
            mock.patch.object(excel_ogrenme, "kayit_onerileri", return_value={}):
        istek.do_POST()
    return istek


def yukle(**k) -> Istek:
    veri = jpeg(len(k.get("cerez") or "") + 3)
    return post("/api/upload", ham=veri,
                basliklar={"Content-Type": "image/jpeg"}, **k)


class TekKullaniciModuTesti(Temel):
    def test_yerel_tam_yetkili_ag_reddedilir(self):
        self.yaz(self.data, [onayli("a.jpg", "100001")])
        with mock.patch.object(excel_ogrenme, "kayit_onerileri", return_value={}):
            self.assertEqual(get("/api/receipts").durum, 200)
            self.assertEqual(get("/api/receipts", istemci=AGDAN).durum, 401)
        self.assertEqual(post("/api/save", {"filename": "a.jpg", "fields": {}},
                              istemci=AGDAN).durum, 401)
        self.assertEqual(get("/", istemci=AGDAN).yonlendirmeler, ["/giris"])
        self.assertEqual(get("/cek", istemci=AGDAN).yonlendirmeler, ["/giris?sonra=/cek"])

    def test_telefon_anahtari_yalniz_yukleme_yetkisi_verir(self):
        self.yaz(self.data, [onayli("a.jpg", "100001")])
        anahtar = erisim.yukleme_anahtari()
        self.assertEqual(get("/cek?anahtar=yanlis", istemci=AGDAN).yonlendirmeler,
                         ["/giris?sonra=/cek"])
        giris = get(f"/cek?anahtar={anahtar}", istemci=AGDAN)
        self.assertEqual(giris.yonlendirmeler, ["/cek"])
        self.assertIn("HttpOnly", giris.cerezler[0])
        cerez = giris.cerezler[0].split(";")[0]

        istek = yukle(istemci=AGDAN, cerez=cerez)
        self.assertEqual(istek.durum, 202, istek.yanitlar)
        yeni = istek.yanitlar[0][1]["filename"]
        self.assertEqual(post("/api/save", {"filename": "a.jpg", "fields": {}},
                              istemci=AGDAN, cerez=cerez).durum, 403)
        self.assertEqual(post("/api/export-excel", {}, istemci=AGDAN, cerez=cerez).durum, 403)
        with mock.patch.object(excel_ogrenme, "kayit_onerileri", return_value={}):
            gorulen = get("/api/receipts", istemci=AGDAN, cerez=cerez).yanitlar[0][1]["receipts"]
        # Telefon defterin tamamını değil yalnız kendi yüklemesini görür.
        self.assertEqual([r["_dosya"] for r in gorulen], [yeni])
        self.assertEqual(self.oku()[1]["_yukleyen"], "telefon")
        self.assertEqual(denetim.oku(yeni)[0]["kullanici"], "telefon")

        # Anahtar yenilenince eski bağlantı çalışmaz.
        erisim.yukleme_anahtari(yenile=True)
        self.assertEqual(yukle(istemci=AGDAN, cerez=cerez).durum, 401)


class KullaniciModuTesti(Temel):
    def setUp(self) -> None:
        super().setUp()
        erisim.kullanici_ekle("ayse", "parola-ayse", "yonetici")
        erisim.kullanici_ekle("veli", "parola-veli", "kontrol")
        erisim.kullanici_ekle("sofor", "parola-sofor", "yukleme")
        self.yaz(self.data, [onayli("a.jpg", "100001", _onaylandi=False)])
        (self.fisler / "a.jpg").write_bytes(jpeg(1))

    def giris(self, ad: str, parola: str | None = None, istemci: str = YEREL) -> str:
        istek = post("/api/giris", {"ad": ad, "parola": parola or f"parola-{ad}"}, istemci=istemci)
        self.assertEqual(istek.durum, 200, istek.yanitlar)
        return istek.cerezler[0].split(";")[0]

    def test_bu_bilgisayar_dahil_giris_zorunlu(self):
        self.assertEqual(get("/api/receipts").durum, 401)
        self.assertEqual(get("/").yonlendirmeler, ["/giris"])
        self.assertEqual(post("/api/giris", {"ad": "veli", "parola": "yanlis"}).durum, 401)
        self.assertEqual(denetim.oku(islem="giris_basarisiz")[0]["kullanici"], "veli")

    def test_rol_sinirlari(self):
        veli = self.giris("veli")
        sofor = self.giris("sofor")
        ayse = self.giris("ayse")
        with mock.patch.object(excel_ogrenme, "kayit_onerileri", return_value={}):
            self.assertEqual(get("/api/receipts", cerez=veli).durum, 200)
        self.assertEqual(post("/api/save", {"filename": "a.jpg", "fields": {}},
                              cerez=sofor).durum, 403)
        self.assertEqual(get("/", cerez=sofor).yonlendirmeler, ["/cek"])
        # Tüm denetim geçmişi yalnız yöneticiye; fiş geçmişi kontrol rolüne açık.
        self.assertEqual(get("/api/denetim", cerez=veli).durum, 403)
        self.assertEqual(get("/api/denetim", cerez=ayse).durum, 200)
        self.assertEqual(get("/api/denetim?dosya=a.jpg", cerez=veli).durum, 200)

    def test_cikis_oturumu_kapatir(self):
        veli = self.giris("veli")
        post("/api/cikis", {}, cerez=veli)
        self.assertEqual(get("/api/durum", cerez=veli).durum, 401)

    def test_kaba_kuvvet_sinirlanir(self):
        for _ in range(erisim.HATALI_GIRIS_SINIRI):
            post("/api/giris", {"ad": "veli", "parola": "yanlis"}, istemci=AGDAN)
        istek = post("/api/giris", {"ad": "veli", "parola": "parola-veli"}, istemci=AGDAN)
        self.assertEqual(istek.durum, 401)
        self.assertIn("Çok fazla", istek.yanitlar[0][1]["error"])
        self.giris("veli", istemci=YEREL)  # başka adres etkilenmez

    def test_bozuk_kullanici_dosyasi_erisimi_acmaz(self):
        erisim.KULLANICI_FILE.write_text("{bozuk", encoding="utf-8")
        self.assertEqual(get("/api/receipts").durum, 401)

    def test_duzenleme_onay_aktarim_ve_silme_kim_ile_kaydedilir(self):
        veli, ayse = self.giris("veli"), self.giris("ayse")
        self.assertEqual(post("/api/save", {"filename": "a.jpg",
                                            "fields": {"litre": "45"}}, cerez=veli).durum, 200)
        self.assertEqual(post("/api/save", {"filename": "a.jpg", "fields": {},
                                            "approved": True}, cerez=ayse).durum, 200)
        kayit = self.oku()[0]
        self.assertEqual((kayit["_son_duzenleyen"], kayit["_onaylayan"]), ("ayse", "ayse"))

        defter = defter_olustur(self.kok / "output" / "defter.xlsx")
        istek = post("/api/export-excel", {"dosya": str(defter), "sayfa": SAYFA}, cerez=veli)
        self.assertEqual(istek.durum, 200, istek.yanitlar)
        kayit = self.oku()[0]
        self.assertEqual(kayit["_excel_aktaran"], "veli")
        self.assertEqual(load_workbook(defter)[SAYFA].cell(row=kayit["_excel_satir"], column=8).value, 45)

        gecmis = denetim.oku(kayit["_yukleme_id"])
        self.assertEqual([(o["islem"], o["kullanici"]) for o in reversed(gecmis)], [
            ("duzenle", "veli"), ("onayla", "ayse"), ("excel_aktar", "veli")])
        self.assertEqual(gecmis[-1]["degisiklikler"], {"litre": ["50", "45"]})
        self.assertEqual(gecmis[0]["satirlar"], [kayit["_excel_satir"]])

        self.assertEqual(post("/api/delete", {"filename": "a.jpg"}, cerez=veli).durum, 200)
        self.assertEqual(self.oku()[0]["_silen"], "veli")
        self.assertEqual(denetim.oku("a.jpg")[0]["islem"], "sil")

    def test_denetim_yazilamazsa_hata_gorunur(self):
        veli = self.giris("veli")
        with mock.patch.object(denetim, "kaydet", side_effect=OSError("disk dolu")):
            istek = post("/api/save", {"filename": "a.jpg", "fields": {}}, cerez=veli)
        self.assertEqual(istek.durum, 400)
        self.assertIn("disk dolu", istek.yanitlar[0][1]["error"])


class TekrarDenemeVeDurumTesti(Temel):
    def test_basarisiz_ocr_gorunur_ve_guvenle_tekrar_denenir(self):
        self.yaz(self.data, [
            {"_dosya": "h.jpg", "_yukleme_id": "u1", "HATA": "Ollama yanıt vermedi",
             "_isleniyor": False},
            {"_dosya": "kayip.jpg", "_yukleme_id": "u2", "HATA": "x", "_isleniyor": False},
        ])
        (self.fisler / "temp").mkdir()
        (self.fisler / "temp" / "h.jpg").write_bytes(jpeg(5))
        durum = get("/api/durum").yanitlar[0][1]
        self.assertEqual(durum["ocr_hatali"], ["h.jpg", "kayip.jpg"])

        istek = post("/api/retry", {"filename": "h.jpg"})
        self.assertEqual(istek.durum, 200, istek.yanitlar)
        kayit = self.oku()[0]
        self.assertTrue(kayit["_isleniyor"])
        self.assertNotIn("HATA", kayit)
        self.assertEqual(kayit["_deneme_sayisi"], 1)
        self.assertEqual(arayuz.OCR_QUEUE.qsize(), 1)
        # Tekrar basmak işi ikiye katlamaz; görseli olmayan da kuyruğa girmez.
        self.assertEqual(post("/api/retry", {"filename": "h.jpg"}).durum, 409)
        self.assertEqual(post("/api/retry", {"filename": "kayip.jpg"}).durum, 404)
        self.assertEqual(arayuz.OCR_QUEUE.qsize(), 1)
        self.assertEqual(denetim.oku("h.jpg")[0]["islem"], "tekrar_dene")

    def test_yarim_excel_aktarimi_durumda_gorunur(self):
        self.yaz(self.data, [onayli("a.jpg", "100001", _excel_bekleyen={
            "islem_id": "x", "dosya": "d.xlsx", "sayfa": SAYFA, "satir": 5})])
        self.assertEqual(get("/api/durum").yanitlar[0][1]["excel_yarim"], ["a.jpg"])
